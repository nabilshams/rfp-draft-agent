"""Building RFP files (Word, Excel, PDF) from the plain-text specs in ``evals/golden/``.

The golden set is stored as readable JSON, not binary files, so it can be reviewed in a pull request and contains
nothing that isn't visible. The files are produced when an evaluation runs.
"""

from __future__ import annotations

import io
from typing import Any


def render_docx(blocks: list[dict[str, Any]]) -> bytes:
    """A Word document from blocks: ``{"p": text}`` paragraphs and ``{"table": [[cell, ...], ...]}`` tables."""
    import docx

    document = docx.Document()
    for block in blocks:
        if "p" in block:
            document.add_paragraph(block["p"])
        elif "table" in block:
            rows = block["table"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            raise ValueError(f"unknown block: {block}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_xlsx(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_text(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def render_pdf(pages: list[list[str]]) -> bytes:
    """A minimal text PDF, one list of lines per page."""
    objects: list[bytes] = []

    def add(body: str) -> int:
        objects.append(body.encode("latin-1"))
        return len(objects)

    add("<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    add(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    font = 3 + 2 * len(pages)
    for i, lines in enumerate(pages):
        content = "BT /F1 12 Tf 14 TL 50 700 Td " + " ".join(f"({_pdf_text(line)}) Tj T*" for line in lines) + " ET"
        add(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>"
        )
        add(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
    add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def render_rfp(spec: dict[str, Any]) -> bytes:
    """The file described by an RFP spec (``format`` is ``docx``, ``xlsx`` or ``pdf``)."""
    kind = spec["format"]
    if kind == "docx":
        return render_docx(spec["blocks"])
    if kind == "xlsx":
        return render_xlsx(spec["rows"])
    if kind == "pdf":
        return render_pdf(spec["pages"])
    raise ValueError(f"unknown RFP format: {kind}")
