from app.outputs.docx_writer import parse_template_sections
from pathlib import Path
from docx import Document

path = Path("templates/rfp_response_template.docx")
print(f"File exists: {path.exists()}, size: {path.stat().st_size}")

doc = Document(str(path))
print(f"Total paragraphs: {len(doc.paragraphs)}")
print()

styles_seen = set()
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        style = p.style.name if p.style else "None"
        styles_seen.add(style)
        if i < 30:
            bold = any(r.bold for r in p.runs) if p.runs else False
            text = p.text[:60]
            print(f"  [{i}] style={style!r:25s} bold={bold!s:5s}  {text}")

print(f"\nUnique styles: {sorted(styles_seen)}")
print()

sections = parse_template_sections(path)
print(f"Detected sections: {len(sections)}")
for s in sections:
    ph = s.get("placeholder", "")
    print(f"  H{s['level']}: {s['heading'][:50]}  placeholder={ph}")
