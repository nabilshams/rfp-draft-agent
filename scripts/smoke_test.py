"""Quick smoke test for local endpoints."""

import httpx
import json
from pathlib import Path

base = "http://localhost:8000"

# 1. Health check
r = httpx.get(f"{base}/health")
print(f"1. Health: {r.json()}")

# 2. Upload a test RFP - create sample inline
from docx import Document as DocxDocument

sample = Path("tests/fixtures/sample_rfp.docx")
sample.parent.mkdir(parents=True, exist_ok=True)
if not sample.exists():
    doc = DocxDocument()
    doc.add_heading("Sample RFP", level=0)
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("This RFP seeks qualified vendors for a cloud-based document management system.")
    doc.add_heading("2. Requirements", level=1)
    doc.add_paragraph("The vendor shall provide a web-based portal accessible 24/7.")
    doc.add_paragraph("The contractor must ensure 99.9% uptime availability.")
    doc.add_paragraph("The provider shall implement role-based access control (RBAC).")
    doc.add_paragraph("The vendor must support integration with Azure Active Directory.")
    doc.add_heading("3. Questions", level=1)
    doc.add_paragraph("Does your platform support multi-tenant architecture?")
    doc.add_paragraph("What is your disaster recovery strategy?")
    doc.save(str(sample))
    print("   Created sample RFP")

with open(sample, "rb") as f:
    r = httpx.post(f"{base}/rfp/upload", files={"file": ("sample_rfp.docx", f)}, timeout=30)
print(f"2. Upload: {r.status_code} - {r.json()}")
rfp_id = r.json()["rfp_id"]

# 3. Extract
r = httpx.post(
    f"{base}/rfp/extract",
    json={"rfp_id": rfp_id, "filename": "sample_rfp.docx"},
    timeout=30,
)
print(f"3. Extract: {r.status_code}")
extraction = r.json()
print(f"   Sections: {len(extraction.get('sections', []))}")
print(f"   Requirements: {len(extraction.get('requirements', []))}")
print(f"   Questions: {len(extraction.get('questions', []))}")
for req in extraction.get("requirements", [])[:3]:
    print(f"   - [{req['priority']}] {req['text'][:80]}")

# 4. Metrics
r = httpx.get(f"{base}/metrics")
print(f"4. Metrics: {r.json()}")

print("\nAll basic endpoints working!")
