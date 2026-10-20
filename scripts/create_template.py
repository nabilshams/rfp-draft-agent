"""Create a minimal DOCX template with placeholders.

Run this script once to generate templates/rfp_response_template.docx.
The template contains placeholder tags that docx_writer will replace.
"""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates"


def create_template():
    doc = Document()

    # Title page
    title = doc.add_heading("RFP Response", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")
    doc.add_paragraph("Prepared by: [Company Name]")
    doc.add_paragraph("Date: [Date]")
    doc.add_page_break()

    # Table of Contents placeholder
    doc.add_heading("Table of Contents", level=1)
    doc.add_paragraph("[Auto-generated table of contents]")
    doc.add_page_break()

    # Executive Summary
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph("{{executive_summary}}")
    doc.add_page_break()

    # Solution Overview
    doc.add_heading("2. Solution Overview", level=1)
    doc.add_paragraph("{{solution_overview}}")
    doc.add_page_break()

    # Compliance Matrix Summary
    doc.add_heading("3. Compliance Matrix Summary", level=1)
    doc.add_paragraph("{{compliance_matrix_summary}}")
    doc.add_page_break()

    # Detailed Responses (filled dynamically)
    doc.add_heading("4. Detailed Responses", level=1)
    doc.add_paragraph("[Detailed responses are appended below by the agent.]")

    # Save
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    output_path = TEMPLATE_DIR / "rfp_response_template.docx"
    doc.save(str(output_path))
    print(f"Template created: {output_path}")


if __name__ == "__main__":
    create_template()
