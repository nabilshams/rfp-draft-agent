"""Demo script – uploads a sample RFP and generates draft outputs.

Usage:
    python scripts/demo.py --base-url http://localhost:8000

For authenticated environments, set the DEMO_TOKEN env var to a valid
Entra ID access token.  For local dev (ENTRA_TENANT_ID=dev), no token
is needed.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import httpx

SAMPLE_RFP_DIR = Path(__file__).resolve().parents[1] / "tests" / "fixtures"


def _headers() -> dict[str, str]:
    token = os.environ.get("DEMO_TOKEN", "")
    if token:
        return {"Authorization": f"Bearer {token}"}
    return {}


def main():
    parser = argparse.ArgumentParser(description="RFP Draft Agent demo")
    parser.add_argument("--base-url", default="http://localhost:8000", help="API base URL")
    args = parser.parse_args()

    base = args.base_url.rstrip("/")
    headers = _headers()

    print("=" * 60)
    print("  RFP Draft Agent – Demo Script")
    print("=" * 60)

    # 1. Health check
    print("\n[1/5] Health check...")
    r = httpx.get(f"{base}/health", timeout=10)
    print(f"  Status: {r.json()}")

    # 2. Upload RFP
    print("\n[2/5] Uploading sample RFP...")
    sample_path = SAMPLE_RFP_DIR / "sample_rfp.docx"
    if not sample_path.exists():
        print(f"  ⚠ Sample RFP not found at {sample_path}")
        print("  Creating a minimal text-based sample...")
        _create_sample_rfp(sample_path)

    with open(sample_path, "rb") as f:
        r = httpx.post(
            f"{base}/rfp/upload",
            files={"file": (sample_path.name, f, "application/octet-stream")},
            headers=headers,
            timeout=30,
        )
    if r.status_code != 200:
        print(f"  Upload failed: {r.status_code} – {r.text}")
        sys.exit(1)

    upload = r.json()
    rfp_id = upload["rfp_id"]
    filename = upload["filename"]
    print(f"  RFP ID: {rfp_id}")
    print(f"  Filename: {filename}")

    # 3. Extract requirements
    print("\n[3/5] Extracting requirements...")
    r = httpx.post(
        f"{base}/rfp/extract",
        json={"rfp_id": rfp_id, "filename": filename},
        headers=headers,
        timeout=30,
    )
    if r.status_code != 200:
        print(f"  Extract failed: {r.status_code} – {r.text}")
        sys.exit(1)

    extraction = r.json()
    print(f"  Sections: {len(extraction.get('sections', []))}")
    print(f"  Requirements: {len(extraction.get('requirements', []))}")
    print(f"  Questions: {len(extraction.get('questions', []))}")

    # 4. Generate draft
    print("\n[4/5] Generating draft response...")
    r = httpx.post(
        f"{base}/draft",
        json={
            "rfp_id": rfp_id,
            "client_id": "demo-client",
            "project_id": "demo-project",
            "solution_area": "cloud",
            "tone": "professional",
        },
        headers=headers,
        timeout=300,
    )
    if r.status_code != 200:
        print(f"  Draft failed: {r.status_code} – {r.text}")
        sys.exit(1)

    draft = r.json()
    print(f"  Elapsed: {draft.get('elapsed_seconds', '?')}s")
    print(f"  DOCX URL: {draft.get('outputs', {}).get('docx_url', 'N/A')}")
    print(f"  XLSX URL: {draft.get('outputs', {}).get('xlsx_url', 'N/A')}")

    responses = draft.get("draft", {}).get("responses", [])
    print(f"  Responses generated: {len(responses)}")

    # 5. Single question
    print("\n[5/5] Asking a single grounded question...")
    r = httpx.post(
        f"{base}/answer",
        json={
            "question": "Does your platform support multi-tenant architecture?",
            "client_id": "demo-client",
            "project_id": "demo-project",
            "solution_area": "cloud",
        },
        headers=headers,
        timeout=120,
    )
    if r.status_code == 200:
        answer = r.json()
        print(f"  Answer: {answer.get('answer_markdown', 'N/A')[:200]}")
    else:
        print(f"  Answer failed: {r.status_code} – {r.text}")

    print("\n" + "=" * 60)
    print("  Demo complete!")
    print("=" * 60)


def _create_sample_rfp(path: Path):
    """Create a minimal DOCX sample RFP for testing."""
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Sample Request for Proposal", level=0)

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "This Request for Proposal (RFP) seeks qualified vendors to provide "
        "a cloud-based document management system."
    )

    doc.add_heading("2. Requirements", level=1)
    doc.add_paragraph("The vendor shall provide a web-based portal accessible 24/7.")
    doc.add_paragraph("The contractor must ensure 99.9% uptime availability.")
    doc.add_paragraph("The provider shall implement role-based access control (RBAC).")
    doc.add_paragraph("The vendor must support integration with Azure Active Directory.")
    doc.add_paragraph("The system shall provide RESTful APIs for all core operations.")

    doc.add_heading("3. Evaluation Criteria", level=1)
    doc.add_paragraph("Technical approach (40%), Experience (30%), Cost (30%).")

    doc.add_heading("4. Questions", level=1)
    doc.add_paragraph("Does your platform support multi-tenant architecture?")
    doc.add_paragraph("What is your disaster recovery strategy?")

    doc.save(str(path))
    print(f"  Created sample RFP: {path}")


if __name__ == "__main__":
    main()
