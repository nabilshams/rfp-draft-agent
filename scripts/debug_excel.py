import os

from azure.identity import DefaultAzureCredential
from azure.storage.blob import BlobServiceClient
from io import BytesIO

blob_service = BlobServiceClient(
    account_url=os.environ["AZURE_STORAGE_ACCOUNT_URL"],
    credential=DefaultAzureCredential(),
)
container = blob_service.get_container_client("rfp-uploads")
data = container.download_blob(
    "f6137a6f-41f/Enterprise_AI_Agent_Platform_Vendor_Capability_Matrix.xlsx"
).readall()
print(f"File size: {len(data)} bytes")
print(f"First 4 bytes (magic): {data[:4].hex()}")

if data[:2] == b"PK":
    print("Format: ZIP-based (xlsx)")
elif data[:8] == bytes.fromhex("D0CF11E0A1B11AE1"):
    print("Format: OLE2 Compound (xls)")
else:
    print(f"Format: Unknown magic {data[:8].hex()}")

# Try openpyxl
print("\n--- openpyxl ---")
try:
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    print(f"Sheet: {ws.title}, max_row={ws.max_row}, max_col={ws.max_column}")
    for i, row in enumerate(ws.iter_rows(values_only=True, max_row=4)):
        cells = [str(c)[:50] if c is not None else "" for c in row]
        print(f"  Row {i}: {cells}")
except Exception as e:
    print(f"FAILED: {e}")

# Try xlrd
print("\n--- xlrd ---")
try:
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    ws = wb.sheet_by_index(0)
    print(f"Sheet: {ws.name}, nrows={ws.nrows}, ncols={ws.ncols}")
    for r in range(min(4, ws.nrows)):
        cells = [str(ws.cell_value(r, c))[:50] for c in range(min(6, ws.ncols))]
        print(f"  Row {r}: {cells}")
except Exception as e:
    print(f"FAILED: {e}")

# Try our extractor
print("\n--- extract_requirements_from_spreadsheet ---")
try:
    from app.ingestion.rfp_extract import extract_requirements_from_spreadsheet
    result = extract_requirements_from_spreadsheet(
        "Enterprise_AI_Agent_Platform_Vendor_Capability_Matrix.xlsx", data
    )
    print(f"Requirements: {len(result.requirements)}")
    print(f"Questions: {len(result.questions)}")
    for r in result.requirements[:3]:
        print(f"  [{r.id}] [{r.priority}] {r.text[:80]}")
except Exception as e:
    import traceback
    traceback.print_exc()
