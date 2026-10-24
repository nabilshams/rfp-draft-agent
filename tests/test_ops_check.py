"""The operator's read-only checks: they report what is there without printing anybody's text."""

import io
from types import SimpleNamespace

import docx

from app.ledger import store as ledger_store
from app.ops import check

from .test_ledger_api import api, api_settings, make_project  # noqa: F401  (api is a fixture)
from .test_ledger_items import REQS

SECRET = "Our secret customer result was 41% last year"


class TestProjectReport:
    async def test_it_summarises_the_items_the_record_and_the_audit_trail(self, api):  # noqa: F811
        project = await make_project(api, requirements=REQS)
        lines = check.project_report(ledger_store.store_for(api_settings()), project["project"]["id"])
        text = "\n".join(lines)
        assert "PROJECT name='Harbourview bid'" in text and "requirement_count=5" in text
        assert "'instruction': 1" in text and "'deadline': 1" in text and "repeats=1" in text and "answered=0" in text
        assert "R3       mandatory" in text and "EVENTS" in text and "project.created" in text

    async def test_it_never_prints_an_answer_and_cuts_a_long_requirement(self, api):  # noqa: F811
        project = await make_project(api, requirements=[{"id": "R1", "text": "The vendor shall " + "x" * 300}])
        from app.ledger import service

        pid = project["project"]["id"]
        store = ledger_store.store_for(api_settings())
        service.add_answer(store, api_settings(), pid, "R1", {"answer_markdown": SECRET}, actor=service.AGENT, source="model")
        text = "\n".join(check.project_report(store, pid))
        assert SECRET not in text and "answered=1" in text
        line = next(x for x in text.splitlines() if x.startswith("  R1"))
        assert len(line.split("| ", 1)[1]) == check.LABEL

    def test_an_unknown_project_is_said_so(self):
        assert check.project_report(ledger_store.store_for(api_settings()), "nope") == ["No project nope"]


class FakeSearch:
    def __init__(self, chunks):
        self.chunks = chunks

    def search(self, search_text="*", filter=None, include_total_count=False, top=None, select=None):  # noqa: A002
        rows = [c for c in self.chunks if self._match(c, filter)]
        result = [{"content": c["content"]} for c in rows][: top or None] if select else []
        return iter(result) if select else SimpleNamespace(get_count=lambda: len(rows))

    @staticmethod
    def _match(chunk, filter):  # noqa: A002
        if not filter:
            return True
        if filter == "page ge 1":
            return (chunk.get("page") or 0) >= 1
        key, _, value = filter.partition(" eq ")
        return chunk.get(key) == value.strip("'")


class TestIndexReport:
    def test_it_counts_the_chunks_that_have_provenance(self, monkeypatch):
        chunks = [{"source_type": "authoritative", "approval_status": "approved", "page": 2, "content": "a"},
                  {"source_type": "historical", "approval_status": "unreviewed", "content": "b"},
                  {"content": "c"}]
        monkeypatch.setattr(check.azure_clients, "search_client", lambda _s: FakeSearch(chunks))
        text = "\n".join(check.index_report(api_settings()))
        assert "3 chunks" in text and "'authoritative': 1" in text and "untyped=1" in text
        assert "none=1" in text and "chunks with a page number=1" in text


class TestOutputsReport:
    def test_figures_the_knowledge_base_does_not_contain_are_listed(self):
        assert check.figures_not_in("We cut costs by 41% and 99.95% uptime.", "Uptime is 99.95% each month.") == ["41%"]
        assert check.figures_not_in("No figures here.", "x") == []

    def test_a_document_is_read_with_its_tables(self):
        document = docx.Document()
        document.add_paragraph("Body 12%")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Cell 34%"
        buffer = io.BytesIO()
        document.save(buffer)
        text = check._docx_text(buffer.getvalue())
        assert "Body 12%" in text and "Cell 34%" in text

    def test_it_lists_the_documents_and_flags_invented_figures(self, monkeypatch):
        document = docx.Document()
        document.add_paragraph("Retention improved by 41% and uptime is 99.95%.")
        buffer = io.BytesIO()
        document.save(buffer)
        blob = SimpleNamespace(name="p1/draft.docx", size=10, last_modified=__import__("datetime").datetime(2026, 1, 2, 3, 4))

        class Container:
            def list_blobs(self, name_starts_with=None):
                return [blob]

            def download_blob(self, name):
                return SimpleNamespace(readall=lambda: buffer.getvalue())

        monkeypatch.setattr(check.azure_clients, "container", lambda _s, _n: Container())
        monkeypatch.setattr(check, "_knowledge_text", lambda _s: "Uptime is 99.95% monthly.")
        text = "\n".join(check.outputs_report(api_settings()))
        assert "p1/draft.docx  10 bytes  2026-01-02 03:04" in text and "figures not in the knowledge base: ['41%']" in text


class TestCommand:
    def test_an_unknown_command_prints_the_help(self, capsys):
        assert check.main(["nonsense"]) == 2
        assert "python -m app.ops.check" in capsys.readouterr().out

    def test_a_project_report_needs_a_ledger(self, monkeypatch, capsys):
        monkeypatch.setattr(check.ledger_store, "store_for", lambda _s: None)
        assert check.main(["project", "abc"]) == 1

class TestEventDetails:
    def test_a_choice_is_shown_as_from_and_to_but_text_is_only_named(self):
        event = {"details": {"changes": {"item_kind": {"from": "instruction", "to": "question"}, "text": {"from": "secret a", "to": "secret b"}}}}
        shown = check._what_changed(event)
        assert "item_kind: 'instruction' -> 'question'" in shown and "text (changed)" in shown
        assert "secret" not in shown

    def test_an_event_without_changes_adds_nothing(self):
        assert check._what_changed({"details": {}}) == "" and check._what_changed({}) == ""

    def test_the_output_is_ascii_whatever_the_console(self, monkeypatch, capsys):
        monkeypatch.setattr(check, "project_report", lambda *_: ["anchor \u00b67"])
        monkeypatch.setattr(check.ledger_store, "store_for", lambda _s: object())
        assert check.main(["project", "abc"]) == 0
        assert "anchor \\xb67" in capsys.readouterr().out
