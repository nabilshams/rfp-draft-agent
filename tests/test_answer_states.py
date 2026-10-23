"""Answer states and the "never invent" policy: what state an answer ends in, and what is kept of the model's text."""

import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook

from app.agents import answers, states, strategies
from app.agents.kb_search import RetrievedContext
from app.config import Settings
from app.ledger import drafting, service
from app.ledger import store as ledger_store
from app.outputs.docx_writer import generate_docx
from app.outputs.xlsx_writer import generate_xlsx

from .test_authorization_api import A
from .test_draft_jobs import worker  # noqa: F401  (a fixture)
from .test_ledger_api import api, api_settings, make_project, url  # noqa: F401  (api is a fixture)
from .test_ledger_drafting import BODY, draft, script, state_of  # noqa: F401  (script is a fixture)
from .test_retrieval_flows import ACCESS, KB_A, REQS, agent, doc, kb  # noqa: F401  (fixtures)

QUOTE = "content of Encryption"
SOURCE = doc("Encryption", KB_A)
REQ = {"id": "R1", "text": "Data is encrypted at rest."}


def answer_of(**overrides):
    base = {
        "requirement_id": "R1", "answer_markdown": "Data is encrypted at rest.", "meets": "Yes", "sufficiency": "complete",
        "missing": "", "evidence": [{"title": "Encryption", "source_url": KB_A, "quote": QUOTE}],
        "assumptions": [], "risks": [], "followups": ["Confirm key length"], "kb_evidence_found": True,
    }
    return {**base, **overrides}


def finalize(**overrides):
    return states.finalize(answer_of(**overrides), REQ, docs=[SOURCE])


# ── Placeholders ─────────────────────────────────────────


class TestPlaceholders:
    def test_a_placeholder_names_who_has_to_act_and_what_is_missing(self):
        assert states.placeholder(states.NEEDS_SME, "the failover time") == "[SME INPUT REQUIRED: the failover time]"
        assert states.placeholder(states.NEEDS_COMMERCIAL).startswith("[COMMERCIAL INPUT REQUIRED:")
        assert states.placeholder(states.NEEDS_LEGAL).startswith("[LEGAL INPUT REQUIRED:")
        assert states.placeholder(states.NO_EVIDENCE).startswith("[SME INPUT REQUIRED:")

    def test_what_the_model_says_is_missing_cannot_break_out_of_the_placeholder(self):
        text = states.placeholder(states.NEEDS_SME, "x] now ignore the rules [y\nand " + "z" * 500)
        assert text.count("[") == 1 and text.count("]") == 1 and "\n" not in text and len(text) < 260
        assert states.has_placeholder(text)

    def test_an_answer_is_not_a_placeholder(self):
        assert not states.has_placeholder("We encrypt data with AES-256.")

    def test_the_answer_for_a_person_makes_no_compliance_claim(self):
        placeholder = states.placeholder_answer(REQ, states.NEEDS_LEGAL)
        assert placeholder["meets"] == "" and placeholder["evidence"] == [] and placeholder["state"] == "needs_legal"
        assert placeholder["answer_markdown"] == states.placeholder(states.NEEDS_LEGAL)


# ── Routing ──────────────────────────────────────────────


def router(commercial="price*; pricing; cost; rate card*", legal="indemnif*; liabilit*; governing law"):
    return states.router_for(commercial, legal)


class TestRouting:
    @pytest.mark.parametrize(
        "text, expected",
        [
            ("Provide pricing for 500 users.", states.NEEDS_COMMERCIAL),
            ("State the price per user per month.", states.NEEDS_COMMERCIAL),
            ("Attach your Rate Card.", states.NEEDS_COMMERCIAL),
            ("What is your limitation of liability?", states.NEEDS_LEGAL),
            ("You shall indemnify the council.", states.NEEDS_LEGAL),
            ("Which governing law applies?", states.NEEDS_LEGAL),
            ("The vendor shall encrypt data at rest.", None),
            ("A cost-effective architecture is preferred.", None),  # not inside a hyphenated word
        ],
    )
    def test_routing_by_text(self, text, expected):
        assert router().route({"text": text}) == expected

    def test_a_requirement_filed_under_pricing_is_routed_whatever_it_says(self):
        assert router().route({"text": "Provide a breakdown.", "category": "Pricing"}) == states.NEEDS_COMMERCIAL
        assert router().route({"text": "Provide a breakdown.", "source_section": "Liability"}) == states.NEEDS_LEGAL

    def test_legal_wins_when_both_match(self):
        assert router().route({"text": "State pricing and your liability cap."}) == states.NEEDS_LEGAL

    def test_the_terms_are_configuration(self):
        custom = states.router_for("roadmap", "")
        assert custom.route({"text": "Share your roadmap."}) == states.NEEDS_COMMERCIAL
        assert custom.route({"text": "Provide pricing."}) is None
        assert states.router_for("", "").route({"text": "Provide pricing and liability terms."}) is None

    def test_the_default_terms_route_what_they_should_and_nothing_in_the_golden_set(self):
        from app.evalkit.golden import load_cases

        default = states.router_from_settings(Settings(_env_file=None))
        assert default.route({"text": "Provide pricing for 500 users."}) == states.NEEDS_COMMERCIAL
        assert default.route({"text": "State your limitation of liability and indemnification terms."}) == states.NEEDS_LEGAL
        routed = {
            r["id"]
            for case in load_cases()
            for r in case.requirements
            if r["id"] not in ("RT-09", "RT-10") and default.route(r)
        }
        assert routed == set()


# ── Deciding the state ───────────────────────────────────


class TestFinalize:
    def test_a_cited_complete_answer_is_fully_answered(self):
        result = finalize()
        assert result["state"] == states.FULLY and result["meets"] == "Yes" and result["missing"] == ""
        assert "sufficiency" not in result and result["answer_markdown"] == "Data is encrypted at rest."

    def test_an_answer_the_model_calls_partial_says_what_is_missing(self):
        result = finalize(sufficiency="partial", meets="Partial", missing="the key length")
        assert result["state"] == states.PARTIAL and result["missing"] == "the key length"
        assert result["answer_markdown"] == "Data is encrypted at rest."  # the grounded prose is kept

    def test_a_partial_compliance_claim_is_a_partial_answer_even_without_the_report(self):
        assert finalize(sufficiency="", meets="Partial")["state"] == states.PARTIAL

    def test_an_answer_that_admits_a_gap_is_partial(self):
        assert finalize(answer_markdown="We encrypt data. Not provided in sources: the key length.")["state"] == states.PARTIAL

    def test_sources_that_disagree_are_reported_as_conflicting(self):
        result = finalize(sufficiency="conflict", missing="the 2022 and 2026 policies differ")
        assert result["state"] == states.CONFLICTING and result["answer_markdown"] == "Data is encrypted at rest."

    @pytest.mark.parametrize("sufficiency, state", [("none", states.NO_EVIDENCE), ("needs_sme", states.NEEDS_SME)])
    def test_when_the_model_says_the_sources_do_not_cover_it_its_prose_is_discarded(self, sufficiency, state):
        result = finalize(sufficiency=sufficiency, answer_markdown="Generic plausible text about encryption.",
                          missing="the encryption algorithm")
        assert result["state"] == state and result["meets"] == "" and result["evidence"] == []
        assert "Generic plausible" not in result["answer_markdown"]
        assert result["answer_markdown"].endswith("the encryption algorithm]")
        assert result["followups"] == ["Confirm key length"]  # what a person should ask is kept

    def test_an_answer_that_cites_nothing_real_is_never_kept(self):
        # The quote is not in the retrieved source: the claim is not grounded
        result = finalize(evidence=[{"title": "Encryption", "source_url": KB_A, "quote": "We encrypt everything twice."}])
        assert result["state"] == states.NO_EVIDENCE and "encrypt data" not in result["answer_markdown"]

    def test_an_answer_that_cites_a_source_that_was_not_retrieved_is_never_kept(self):
        result = finalize(evidence=[{"title": "X", "source_url": "https://other/x", "quote": QUOTE}])
        assert result["state"] == states.NO_EVIDENCE

    def test_an_answer_with_no_evidence_at_all_is_never_kept(self):
        assert finalize(evidence=[])["state"] == states.NO_EVIDENCE

    def test_an_answer_that_opens_by_saying_the_sources_do_not_cover_it_is_a_placeholder(self):
        result = finalize(answer_markdown="Not provided in sources – recommend human review.")
        assert result["state"] == states.NO_EVIDENCE and result["meets"] == ""

    def test_web_grounding_stands_in_for_a_quote(self):
        result = states.finalize(answer_of(evidence=[]), REQ, docs=[], web_grounded=True)
        assert result["state"] == states.FULLY

    def test_an_empty_answer_is_a_failure_not_an_answer(self):
        result = finalize(answer_markdown="  ")
        assert result["state"] == states.FAILED and result["generation_failed"] is True and result["meets"] == ""

    def test_a_failure_makes_no_compliance_claim(self):
        failed = states.failed_answer(REQ, message="Error: boom", risk="boom")
        assert failed["state"] == "failed" and failed["meets"] == "" and failed["generation_failed"] is True
        assert failed["risks"] == ["boom"] and failed["evidence"] == []


class TestPreGate:
    def gate(self, requirement=REQ, found=True, web=False):
        return states.pre_gate(requirement, router=router(), evidence_found=found, web_search=web)

    def test_a_requirement_with_evidence_goes_to_the_model(self):
        assert self.gate() is None

    def test_a_routed_requirement_never_does(self):
        gated = self.gate({"id": "P1", "text": "Provide pricing."})
        assert gated["state"] == states.NEEDS_COMMERCIAL and gated["kb_evidence_found"] is True

    def test_nothing_retrieved_means_nothing_to_draft_from(self):
        gated = self.gate(found=False)
        assert gated["state"] == states.NO_EVIDENCE and gated["kb_evidence_found"] is False

    def test_web_grounding_is_a_source_when_it_is_on(self):
        assert self.gate(found=False, web=True) is None


# ── In the drafting pipeline ─────────────────────────────


def stream(requirements, **kwargs):
    return list(strategies.stream_per_requirement(requirements=requirements, tone="professional", **ACCESS, **kwargs))


class TestPipeline:
    def test_a_pricing_requirement_is_not_sent_to_the_model(self, agent, kb):  # noqa: F811
        kb.docs = {"Provide pricing for 500 users.": [SOURCE]}  # even with documents on hand

        (result,) = stream([{"id": "P1", "text": "Provide pricing for 500 users."}])

        assert agent.prompts == []
        assert result["state"] == "needs_commercial" and "COMMERCIAL INPUT REQUIRED" in result["answer_markdown"]
        assert result["meets"] == ""

    def test_a_requirement_with_nothing_in_the_knowledge_base_is_not_sent_to_the_model(self, agent, kb):  # noqa: F811
        (result,) = stream([{"id": "R9", "text": "Are you ISO 9001 certified?"}])
        assert agent.prompts == [] and result["state"] == "no_evidence"

    def test_with_web_search_on_a_requirement_with_no_documents_is_still_drafted(self, agent, kb):  # noqa: F811
        agent.annotation_urls = {"https://azure.microsoft.com/support/legal/sla/"}
        agent.respond = lambda _p: json.dumps(answer_of(evidence=[]))
        (result,) = stream([{"id": "R9", "text": "What is the SLA?"}], allow_web_search=True)
        assert len(agent.prompts) == 1 and result["state"] == "fully_answered"

    def test_the_model_is_told_to_report_sufficiency_and_what_is_missing(self, agent, kb):  # noqa: F811
        kb.docs = {REQS[0]["text"]: [SOURCE]}
        agent.respond = lambda _p: json.dumps(answer_of())
        stream(REQS[:1])
        assert '"sufficiency"' in agent.prompts[0] and '"missing"' in agent.prompts[0]

    def test_the_model_reporting_no_coverage_gives_a_placeholder_without_its_prose(self, agent, kb):  # noqa: F811
        kb.docs = {REQS[0]["text"]: [SOURCE]}
        agent.respond = lambda _p: json.dumps(
            answer_of(sufficiency="none", answer_markdown="We are the best.", evidence=[], missing="a reference")
        )
        (result,) = stream(REQS[:1])
        assert result["state"] == "no_evidence" and "We are the best" not in result["answer_markdown"]

    def test_a_forced_generation_failure_is_recorded_as_failed_not_as_no(self, agent, kb, monkeypatch):  # noqa: F811
        kb.docs = {REQS[0]["text"]: [SOURCE]}

        def boom(*_a, **_k):
            raise RuntimeError("model is down")

        monkeypatch.setattr(answers.runner, "run_agent_prompt", boom)
        (result,) = stream(REQS[:1])
        assert result["state"] == "failed" and result["generation_failed"] is True
        assert result["meets"] == "" and result["evidence"] == []

    def test_a_reply_that_is_not_an_answer_is_a_failure_and_its_text_is_dropped(self, agent, kb):  # noqa: F811
        kb.docs = {REQS[0]["text"]: [SOURCE]}
        agent.respond = lambda _p: "We are fully compliant with everything."
        (result,) = stream(REQS[:1])
        assert result["state"] == "failed" and "fully compliant" not in result["answer_markdown"]

    def test_a_whole_draft_settles_every_response(self, agent, kb):  # noqa: F811
        kb.docs = {"data is encrypted at rest": [SOURCE]}
        requirements = [
            {"id": "R1", "text": "data is encrypted at rest"},
            {"id": "R2", "text": "Provide pricing for 500 users."},
            {"id": "R3", "text": "Prose with no citation"},
        ]
        agent.respond = lambda _p: json.dumps(
            {
                "executive_summary": "s", "solution_overview": "o",
                "responses": [
                    answer_of(requirement_id="R1"),
                    answer_of(requirement_id="R2", answer_markdown="Our prices are low."),
                    answer_of(requirement_id="R3", evidence=[], answer_markdown="We do everything."),
                ],
            }
        )

        draft = answers.run_draft(
            rfp_id="r", requirements_json=json.dumps(requirements), tone="professional", **ACCESS
        )

        by_id = {r["requirement_id"]: r for r in draft["responses"]}
        assert by_id["R1"]["state"] == "fully_answered"
        assert by_id["R2"]["state"] == "needs_commercial" and "Our prices" not in by_id["R2"]["answer_markdown"]
        assert by_id["R3"]["state"] == "no_evidence" and "We do everything" not in by_id["R3"]["answer_markdown"]


class TestNarratives:
    ROUTES = {"Classify each requirement": json.dumps({"Overview": ["R1", "R2"]})}

    def respond(self, prompt):
        if "Classify each requirement" in prompt:
            return json.dumps({"Overview": ["R1", "R2"]})
        if "Respond to this single RFP requirement" in prompt:
            return json.dumps(answer_of(requirement_id="R1" if "Requirement ID: R1" in prompt else "R2"))
        return json.dumps({"section_name": "Overview", "narrative_markdown": "Narrative.", "requirement_ids": []})

    @pytest.mark.parametrize("approach", [strategies.stream_per_section, strategies.stream_map_reduce])
    def test_a_requirement_for_a_person_is_left_out_of_the_narrative(self, agent, kb, approach):  # noqa: F811
        kb.docs = {"data is encrypted at rest": [SOURCE]}
        agent.respond = self.respond
        requirements = [
            {"id": "R1", "text": "data is encrypted at rest"},
            {"id": "R2", "text": "Provide pricing for 500 users."},
        ]

        events = list(approach(requirements=requirements, template_sections=[{"heading": "Overview"}],
                               tone="professional", **ACCESS))

        narrative_prompts = [p for p in agent.prompts if "section of" in p]
        assert narrative_prompts and all("pricing" not in p for p in narrative_prompts)
        narrative = next(e["data"] for e in events if e["type"] == "section_narrative")
        assert narrative["requirement_ids"] == ["R1"]

    def test_the_section_approach_still_reports_the_state_of_each_requirement(self, agent, kb):  # noqa: F811
        kb.docs = {"data is encrypted at rest": [SOURCE]}
        agent.respond = self.respond
        events = list(strategies.stream_per_section(
            requirements=[{"id": "R1", "text": "data is encrypted at rest"}, {"id": "R2", "text": "Provide pricing."}],
            template_sections=[{"heading": "Overview"}], tone="professional", **ACCESS))
        compliance = {e["data"]["requirement_id"]: e["data"] for e in events if e["type"] == "requirement_compliance"}
        assert compliance["R1"]["state"] == "fully_answered" and compliance["R2"]["state"] == "needs_commercial"
        assert compliance["R2"]["meets"] == ""

    def test_a_failure_in_the_section_approach_is_a_failed_state(self, agent, kb, monkeypatch):  # noqa: F811
        monkeypatch.setattr(strategies.sections, "classify_sections", lambda *a, **k: {})
        monkeypatch.setattr(answers, "answer_single_requirement", lambda **k: (_ for _ in ()).throw(RuntimeError("x")))
        events = list(strategies.stream_per_section(
            requirements=[REQS[0]], template_sections=[{"heading": "Overview"}], tone="professional", **ACCESS))
        (compliance,) = [e["data"] for e in events if e["type"] == "requirement_compliance"]
        assert compliance["state"] == "failed" and compliance["meets"] == "" and compliance["generation_failed"] is True


# ── In the ledger ────────────────────────────────────────


def placeholder_reply(requirement, state="no_evidence", missing="a reference"):
    return states.placeholder_answer(requirement, state, missing=missing)


class TestLedger:
    async def test_a_model_answer_keeps_its_state_and_what_is_missing(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: {**answer_of(requirement_id="R1", sufficiency="", meets="Partial"),
                                        "state": "partially_answered", "missing": "the key length"}
        project = await make_project(api)
        await draft(api, worker, project)

        first = (await state_of(api, project))["requirements"][0]
        assert (first["state"], first["missing"], first["status"]) == ("partially_answered", "the key length", "drafted")
        assert first["answer"]["state"] == "partially_answered"

    async def test_a_placeholder_needs_input_and_is_counted_so(self, api, worker, script):  # noqa: F811
        script.hooks["R2"] = lambda r: placeholder_reply({"id": "R2", "text": r["text"]})
        script.hooks["R3"] = lambda r: placeholder_reply({"id": "R3", "text": r["text"]}, "needs_legal", "")
        project = await make_project(api)
        await draft(api, worker, project)

        state = await state_of(api, project)
        by_rid = {r["rid"]: r for r in state["requirements"]}
        assert by_rid["R2"]["status"] == "needs_input" and by_rid["R2"]["state"] == "no_evidence"
        assert by_rid["R2"]["answer"]["meets"] == ""
        assert by_rid["R3"]["state"] == "needs_legal"
        assert state["counts"]["needs_input"] == 2 and state["counts"]["drafted"] == 1
        # a placeholder is an answer the next draft does not redo, but the status shows it still needs a person
        status = (await api.get(url(project, "draft"), headers=A)).json()
        assert status["needs_input"] == 2 and status["remaining"] == 0 and status["failed"] == 0

    async def test_a_placeholder_cannot_be_approved(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: placeholder_reply({"id": "R1", "text": r["text"]}, "needs_commercial", "")
        project = await make_project(api)
        await draft(api, worker, project)

        refused = await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)

        assert refused.status_code == 409 and "needs commercial" in refused.json()["detail"].lower()
        assert (await api.get(url(project, "requirements/R1"), headers=A)).json()["status"] == "needs_input"

    async def test_conflicting_evidence_cannot_be_approved_either(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: {**answer_of(requirement_id="R1"), "state": "conflicting_evidence"}
        project = await make_project(api)
        await draft(api, worker, project)
        assert (await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)).status_code == 409

    async def test_partial_and_full_answers_can_be_approved(self, api, worker, script):  # noqa: F811
        script.hooks["R2"] = lambda r: {**answer_of(requirement_id="R2"), "state": "partially_answered"}
        project = await make_project(api)
        await draft(api, worker, project)
        for rid in ("R1", "R2"):
            approved = await api.patch(url(project, f"requirements/{rid}"), json={"status": "approved"}, headers=A)
            assert approved.status_code == 200 and approved.json()["status"] == "approved"

    async def test_a_person_filling_in_a_placeholder_makes_it_a_full_answer_they_can_approve(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: placeholder_reply({"id": "R1", "text": r["text"]})
        project = await make_project(api)
        await draft(api, worker, project)

        written = await api.post(url(project, "requirements/R1/answers"),
                                 json={"answer_markdown": "We encrypt data at rest with AES-256."}, headers=A)

        body = written.json()
        assert (body["status"], body["state"], body["missing"]) == ("edited", "fully_answered", "")
        approved = await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)
        assert approved.status_code == 200

    async def test_restoring_an_earlier_version_brings_its_state_back(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: placeholder_reply({"id": "R1", "text": r["text"]})
        project = await make_project(api)
        await draft(api, worker, project)
        await api.post(url(project, "requirements/R1/answers"), json={"answer_markdown": "Written."}, headers=A)

        restored = await api.post(url(project, "requirements/R1/answers"), json={"restore_version": 1}, headers=A)

        assert restored.json()["state"] == "no_evidence" and restored.json()["status"] == "needs_input"

    async def test_a_failed_generation_is_marked_failed_keeps_no_answer_and_is_tried_again(self, api, worker, script):  # noqa: F811
        script.hooks["R2"] = lambda r: states.failed_answer({"id": "R2", "text": r["text"]}, message="Error", risk="x")
        project = await make_project(api)
        await draft(api, worker, project)

        state = await state_of(api, project)
        failed = {r["rid"]: r for r in state["requirements"]}["R2"]
        assert failed["state"] == "failed" and failed["answer"] is None and failed["answer_version"] == 0
        status = (await api.get(url(project, "draft"), headers=A)).json()
        assert status["failed"] == 1 and status["remaining"] == 1
        events = (await api.get(url(project, "events"), params={"limit": 50}, headers=A)).json()["events"]
        assert any(e["action"] == "answer.failed" and e["rid"] == "R2" for e in events)

        script.hooks.pop("R2")
        await draft(api, worker, project)  # the next run drafts only R2, and a good answer clears the failure
        again = {r["rid"]: r for r in (await state_of(api, project))["requirements"]}["R2"]
        assert script.drafted()[-1] == "R2" and again["state"] == "fully_answered" and again["answer_version"] == 1

    async def test_the_documents_include_a_failed_requirement(self, api, worker, script):  # noqa: F811
        script.hooks["R2"] = lambda r: states.failed_answer({"id": "R2", "text": r["text"]}, message="Error", risk="x")
        project = await make_project(api)
        await draft(api, worker, project)

        sink = drafting.ProjectSink(
            ledger_store.store_for(api_settings()), api_settings(), project["project"]["id"], "j",
            drafting.select(ledger_store.store_for(api_settings()), project["project"]["id"], ["R1", "R2", "R3"]),
        )
        responses, requirements = sink.document_inputs()

        by_id = {r["requirement_id"]: r for r in responses}
        assert set(by_id) == {"R1", "R2", "R3"} and len(requirements) == 3
        assert by_id["R2"]["state"] == "failed" and by_id["R2"]["meets"] == ""
        assert by_id["R1"]["state"] == "fully_answered"

    async def test_an_old_answer_without_a_state_can_still_be_approved(self, api):  # noqa: F811
        project = await make_project(api)
        await api.post(url(project, "requirements/R1/answers"), json={"answer_markdown": "Old."}, headers=A)
        stored = ledger_store.store_for(api_settings())
        document = stored.get(project["project"]["id"], service.req_doc_id("R1"))
        document.pop("state")
        stored.apply(project["project"]["id"], [ledger_store.replace(document)])

        approved = await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)
        assert approved.status_code == 200

    async def test_a_status_of_needs_input_cannot_be_chosen_by_a_user(self, api):  # noqa: F811
        project = await make_project(api)
        refused = await api.patch(url(project, "requirements/R1"), json={"status": "needs_input"}, headers=A)
        assert refused.status_code == 422


# ── In the documents ─────────────────────────────────────


RESPONSES = [
    {"requirement_id": "R1", "answer_markdown": "AES-256.", "meets": "Yes", "state": "fully_answered", "evidence": [],
     "missing": ""},
    {"requirement_id": "R2", "answer_markdown": "We run in two regions.", "meets": "Partial", "missing": "failover time",
     "state": "partially_answered", "evidence": []},
    {"requirement_id": "R3", "answer_markdown": "[COMMERCIAL INPUT REQUIRED: x]", "meets": "", "state": "needs_commercial",
     "evidence": []},
    {"requirement_id": "R4", "answer_markdown": "Error", "meets": "", "state": "failed", "evidence": []},
    {"requirement_id": "R5", "answer_markdown": "Legacy answer", "meets": "Yes", "evidence": []},
]


class TestDocuments:
    def sheet(self):
        workbook = load_workbook(io.BytesIO(generate_xlsx(RESPONSES, [])))
        return workbook.active

    def test_the_matrix_has_an_answer_state_column_after_the_existing_ones(self):
        sheet = self.sheet()
        assert [sheet.cell(1, c).value for c in (1, 2, 3, 4, 5, 6, 7)][3:] == [
            "Meets (Yes/Partial/No)", "Evidence / Notes", "Source URL", "Answer state"]
        assert [sheet.cell(r, 7).value for r in range(2, 7)] == [
            "Fully answered", "Partially answered", "Needs commercial", "Failed", None]

    def test_a_placeholder_leaves_the_compliance_cell_empty(self):
        sheet = self.sheet()
        assert sheet.cell(4, 4).value in (None, "") and sheet.cell(5, 4).value in (None, "")
        assert sheet.cell(3, 4).value == "Partial"

    def test_what_is_missing_is_in_the_notes(self):
        assert "Missing: failover time" in self.sheet().cell(3, 5).value

    def paragraphs(self):
        document = Document(io.BytesIO(generate_docx({"executive_summary": "s", "solution_overview": "o", "responses": RESPONSES})))
        return [p.text for p in document.paragraphs]

    def test_the_document_shows_each_requirements_status(self):
        text = self.paragraphs()
        assert "Compliance: Yes   Status: Fully answered" in text
        assert "Compliance: Partial   Status: Partially answered" in text
        assert "Missing: failover time" in text

    def test_a_placeholder_is_not_reported_as_compliant_or_non_compliant(self):
        text = self.paragraphs()
        assert "Compliance: Not assessed   Status: Needs commercial" in text
        assert "Compliance: Not assessed   Status: Failed" in text
        assert "Compliance: Yes" in text and not any(t.startswith("Compliance: No ") for t in text)

    def test_an_answer_without_a_state_is_shown_as_before(self):
        assert "Compliance: Yes" in self.paragraphs()


class TestRecordedInTelemetry:
    def test_the_state_may_be_attached_to_a_span_and_is_only_ever_a_name(self):
        from app.telemetry import spans

        kept = spans.safe_attributes({"answer.state": "needs_legal", "answer.text": "secret"})
        assert kept == {"answer.state": "needs_legal"}


def test_the_state_groups_are_consistent():
    assert set(states.STATES) == set(states.LABELS)
    assert set(states.APPROVABLE) <= set(states.ANSWERED)
    assert not set(states.PLACEHOLDER_STATES) & set(states.ANSWERED)
    assert states.FAILED in states.NEEDS_PERSON and states.FAILED not in states.PLACEHOLDER_STATES
    assert RetrievedContext([SOURCE]).found


def test_a_model_answer_without_a_report_is_still_judged_by_its_citations():
    plain = {k: v for k, v in answer_of().items() if k not in ("sufficiency", "missing")}
    assert states.finalize(plain, REQ, docs=[SOURCE])["state"] == states.FULLY
    assert service.STATUSES[-1] == "needs_input"
