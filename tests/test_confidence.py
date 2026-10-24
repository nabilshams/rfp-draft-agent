"""The confidence band: deterministic, explainable, and calibrated against what reviewers did (#36)."""

import copy
import io
import json
from datetime import date

import pytest
from docx import Document
from openpyxl import load_workbook

from app.agents import confidence, states, strategies, verify
from app.config import Settings
from app.evalkit import metrics, report
from app.evalkit.golden import GOLDEN_DIR
from app.evalkit.run import run_golden
from app.ledger import service
from app.outputs.docx_writer import generate_docx
from app.outputs.xlsx_writer import generate_xlsx

from .test_answer_states import answer_of
from .test_authorization_api import A, B
from .test_draft_jobs import worker  # noqa: F401  (a fixture)
from .test_ledger_api import api, api_settings, make_project, url  # noqa: F401  (api is a fixture)
from .test_ledger_drafting import draft, script, state_of  # noqa: F401  (script is a fixture)
from .test_retrieval_flows import ACCESS, KB_A, agent, doc, kb  # noqa: F401  (fixtures)
from .test_ui_projects import answer, block, render, requirement

TODAY = date(2026, 10, 1)
SETTINGS = Settings(_env_file=None)


def source(n, *, kind="authoritative", status="approved", effective="2026-05-01", found=True, **extra):
    return {
        "id": f"E{n}", "title": f"Doc {n}", "source_url": f"/kb/doc/doc{n}", "quote": f"quote {n}", "source_type": kind,
        "approval_status": status, "effective_date": effective, "quote_verified": found, **extra,
    }


def claim(n, *ids, verdict="supported"):
    return {"id": f"C{n}", "text": f"claim {n}", "evidence_ids": list(ids), "verdict": verdict}


def summary(supported=2, partial=0, unsupported=0, contradicted=0, status=None):
    total = supported + partial + unsupported + contradicted
    status = status or (
        verify.FLAGGED if unsupported or contradicted else verify.VERIFIED if partial == 0 and total else verify.PARTIALLY_VERIFIED
    )
    return {"status": status, "claims": total, "supported": supported, "partial": partial, "unsupported": unsupported,
            "contradicted": contradicted, "unchecked": 0, "by_rule": 0}


def good(**overrides):
    """An answer that deserves High: two authoritative, current sources, every claim supported."""
    base = {
        "state": states.FULLY, "missing": "", "evidence": [source(1), source(2)],
        "claims": [claim(1, "E1"), claim(2, "E2")], "verification": summary(2),
    }
    return {**base, **overrides}


DOCS = [{"source_url": "/kb/doc/doc1", "rank_score": 1.0}, {"source_url": "/kb/doc/doc2", "rank_score": 0.8},
        {"source_url": "/kb/doc/doc3", "rank_score": 0.4}, {"source_url": "/kb/doc/doc4", "rank_score": 0.3},
        {"source_url": "/kb/doc/doc5", "rank_score": 0.2}]


def assess(answer, docs=DOCS, settings=SETTINGS):
    return confidence.assess(answer, docs, settings=settings, today=TODAY)


def codes(result, effect="lowers"):
    return [r["code"] for r in result["reasons"] if r["effect"] == effect]


# ── Every answer gets a band and the reasons ─────────────


class TestBandAndReasons:
    def test_a_well_founded_answer_is_high_and_says_why(self):
        result = assess(good())
        assert result["band"] == "high" and codes(result) == []
        assert set(codes(result, "supports")) >= {"corroborated", "authoritative", "fresh", "verified", "good_match"}
        assert all(r["text"] for r in result["reasons"])

    def test_the_reasons_that_lower_the_band_come_first_worst_first(self):
        weak = good(evidence=[source(1, kind="historical", effective="2020-01-01")], claims=[claim(1, "E1")],
                    verification=summary(1, unsupported=1), state=states.PARTIAL)
        result = assess(weak)
        levels = [r["level"] for r in result["reasons"] if r["effect"] == "lowers"]
        assert levels == sorted(levels, key=lambda lv: confidence._ORDER[lv]) and levels[0] == "low"
        assert result["reasons"][-1]["effect"] == "lowers" or result["band"] == "low"

    def test_the_band_is_the_lowest_any_reason_allows(self):
        assert assess(good(state=states.PARTIAL))["band"] == "medium"
        assert assess(good(state=states.PARTIAL, verification=summary(1, unsupported=1)))["band"] == "low"

    def test_the_facts_the_band_is_made_from_are_returned_with_it(self):
        features = assess(good())["features"]
        assert features["sources"] == 2 and features["supported_fraction"] == 1.0 and features["best_cited_rank"] == 1
        assert features["corroborated_claims"] == 0 and features["oldest_claim_age_days"] == 153

    def test_describe_reads_the_band_in_a_line(self):
        assert confidence.describe(assess(good())) == "High: " + "; ".join(r["text"] for r in assess(good())["reasons"][:3])
        low = assess(good(state=states.PARTIAL, verification=summary(1, unsupported=1)))
        assert confidence.describe(low).startswith("Low: 1 claim(s) are not supported")
        assert confidence.describe(None) == "" and confidence.describe({"band": "nonsense"}) == ""
        assert confidence.describe({"band": "low"}) == "Low"


# ── Changing one input changes the band predictably ──────


class TestOneInputAtATime:
    def test_removing_a_source_changes_the_band_as_the_rules_say(self):
        # Two authoritative sources: High. Each claim keeps its own source. Remove the second and its claim has none.
        full = good()
        assert assess(full)["band"] == "high"
        without = copy.deepcopy(full)
        without["evidence"] = without["evidence"][:1]
        without["claims"][1]["evidence_ids"] = []
        without["verification"] = summary(1, unsupported=1)  # the verifier flags a claim with nothing behind it
        assert assess(without)["band"] == "low" and "claims_unsupported" in codes(assess(without))

    def test_removing_the_only_source_leaves_nothing_to_trust(self):
        gone = good(evidence=[], claims=[])
        assert assess(gone)["band"] == "low" and "no_verified_source" in codes(assess(gone))

    def test_a_single_authoritative_source_is_enough(self):
        one = good(evidence=[source(1)], claims=[claim(1, "E1")], verification=summary(1))
        assert assess(one)["band"] == "high"

    def test_a_single_unreviewed_source_is_not(self):
        one = good(evidence=[source(1, kind="approved_answer", status="unreviewed")], claims=[claim(1, "E1")], verification=summary(1))
        result = assess(one)
        assert result["band"] == "medium" and codes(result) == ["single_source"]

    def test_a_second_source_corroborates_a_weaker_one(self):
        both = good(evidence=[source(1, kind="approved_answer", status="unreviewed"),
                              source(2, kind="approved_answer", status="unreviewed")],
                    claims=[claim(1, "E1", "E2")], verification=summary(1))
        assert assess(both)["band"] == "high"  # two independent sources behind the claim

    def test_two_chunks_of_one_document_are_one_source(self):
        same = good(evidence=[source(1, kind="approved_answer", status="unreviewed"),
                              {**source(2, kind="approved_answer", status="unreviewed"), "source_url": "/kb/doc/doc1/"}],
                    claims=[claim(1, "E1", "E2")], verification=summary(1))
        result = assess(same)
        assert result["band"] == "medium" and result["features"]["sources"] == 1 and "single_source" in codes(result)

    def test_an_approved_past_answer_can_stand_alone_but_a_superseded_one_cannot(self):
        approved = good(evidence=[source(1, kind="approved_answer")], claims=[claim(1, "E1")], verification=summary(1))
        assert assess(approved)["band"] == "high"
        superseded = good(evidence=[source(1, kind="approved_answer", status="superseded")], claims=[claim(1, "E1")],
                          verification=summary(1))
        assert assess(superseded)["band"] == "medium" and "weak_sources" in codes(assess(superseded))

    def test_historical_material_caps_the_band_at_medium(self):
        old = good(evidence=[source(1, kind="historical", effective="2025-06-01"), source(2, kind="historical", effective="2025-06-01")],
                   claims=[claim(1, "E1", "E2")], verification=summary(1))
        result = assess(old)
        assert result["band"] == "medium" and codes(result) == ["weak_sources"]

    @pytest.mark.parametrize("effective, band, code", [
        ("2026-05-01", "high", None), ("2024-12-01", "high", None),  # under two years
        ("2023-01-01", "medium", "dated"), ("2020-01-01", "low", "stale"),
    ])
    def test_age_moves_the_band_at_the_configured_thresholds(self, effective, band, code):
        aged = good(evidence=[source(1, effective=effective)], claims=[claim(1, "E1")], verification=summary(1))
        result = assess(aged)
        assert result["band"] == band and (code in codes(result) if code else "dated" not in codes(result))

    def test_the_thresholds_are_configuration(self):
        aged = good(evidence=[source(1, effective="2023-01-01")], claims=[claim(1, "E1")], verification=summary(1))
        lenient = Settings(_env_file=None, confidence_fresh_days=2000)
        assert assess(aged, settings=lenient)["band"] == "high"
        strict = Settings(_env_file=None, confidence_stale_days=900)
        assert assess(aged, settings=strict)["band"] == "low"

    def test_a_claim_resting_on_an_old_source_lowers_the_band_even_beside_a_new_one(self):
        # This is the case a whole-answer "newest source" would hide
        mixed = good(evidence=[source(1, effective="2020-01-01", kind="historical"), source(2)],
                     claims=[claim(1, "E1"), claim(2, "E2")], verification=summary(2))
        result = assess(mixed)
        assert result["band"] == "low" and {"stale", "weak_sources"} <= set(codes(result))
        assert "corroborated" in codes(result, "supports")  # still shown: two sources were cited

    def test_a_date_that_is_not_known_is_a_reason_not_a_crash(self):
        undated = good(evidence=[source(1, effective="")], claims=[claim(1, "E1")], verification=summary(1))
        assert "undated" in codes(assess(undated)) and assess(undated)["band"] == "medium"

    def test_the_share_of_supported_claims_sets_the_band(self):
        assert assess(good(verification=summary(9, partial=1)))["band"] == "high"  # 90%
        mid = assess(good(verification=summary(3, partial=2)))
        assert mid["band"] == "medium" and codes(mid) == ["claims_partly_supported"]
        low = assess(good(verification=summary(1, partial=2)))
        assert low["band"] == "low" and "few_claims_supported" in codes(low)

    def test_a_claim_the_evidence_does_not_support_makes_it_low(self):
        assert assess(good(verification=summary(3, unsupported=1)))["band"] == "low"

    def test_a_contradiction_or_a_conflict_makes_it_low(self):
        assert assess(good(verification=summary(1, contradicted=1)))["band"] == "low"
        conflicting = assess(good(state=states.CONFLICTING))
        assert conflicting["band"] == "low" and "conflict" in codes(conflicting)

    def test_a_partial_answer_is_medium_at_best(self):
        result = assess(good(state=states.PARTIAL, missing="the key length"))
        assert result["band"] == "medium" and codes(result) == ["partial"]

    def test_claims_that_were_not_checked_cap_it_at_medium(self):
        for status in ("", verify.UNCHECKED):
            result = assess(good(verification={"status": status} if status else {}, claims=[]))
            assert result["band"] == "medium" and "unchecked" in codes(result)

    def test_a_quote_that_is_not_in_its_source_is_a_reason(self):
        result = assess(good(evidence=[source(1), source(2, found=False)], claims=[claim(1, "E1")], verification=summary(1)))
        assert result["band"] == "medium" and codes(result) == ["quote_not_found"]

    def test_web_sources_cannot_be_checked_so_the_band_stops_at_medium(self):
        web = good(evidence=[source(1, found=None)], claims=[], verification={"status": "unchecked"})
        result = assess(web)
        assert result["band"] == "medium" and "web_only" in codes(result)

    def test_where_the_cited_source_ranked_among_the_matches_matters(self):
        low_rank = good(evidence=[source(4)], claims=[claim(1, "E4")], verification=summary(1))
        result = assess(low_rank)
        assert result["features"]["best_cited_rank"] == 4 and "weak_match" in codes(result) and result["band"] == "medium"

    def test_a_cited_source_that_scored_far_below_the_best_match_is_a_weak_match(self):
        docs = [{"source_url": "/kb/doc/x", "rank_score": 1.0}, {"source_url": "/kb/doc/doc1", "rank_score": 0.2}]
        result = assess(good(evidence=[source(1)], claims=[claim(1, "E1")], verification=summary(1)), docs)
        assert result["features"]["best_cited_relative_score"] == 0.2 and "weak_match" in codes(result)

    def test_without_retrieval_scores_the_match_is_simply_not_judged(self):
        result = assess(good(), docs=[])
        assert result["band"] == "high" and result["features"]["best_cited_rank"] is None

    def test_with_no_claims_listed_the_answer_is_judged_as_one(self):
        result = assess(good(claims=[], verification={"status": "verified", "claims": 0, "supported": 0}))
        assert result["features"]["sources"] == 2 and result["band"] == "high"


# ── No model number is an input ──────────────────────────


class TestNothingTheModelSaysIsAnInput:
    def test_the_models_own_claims_about_itself_change_nothing(self):
        base = assess(good())
        boasting = assess(good(meets="Yes", sufficiency="complete", confidence_score=0.99, model_confidence="very high",
                               assumptions=["I am certain"], answer_markdown="Definitely, 100% sure."))
        assert boasting["band"] == base["band"] and boasting["reasons"] == base["reasons"]
        assert boasting["features"] == base["features"]

    def test_the_inputs_are_only_what_the_system_can_check(self):
        assert set(assess(good())["features"]) == {
            "sources", "unfound_quotes", "unverifiable_sources", "authoritative", "reviewed_answer", "claims_on_weak_sources",
            "claims_on_one_source", "corroborated_claims", "oldest_claim_age_days", "undated_claims", "verification",
            "claims", "supported_fraction", "flagged_claims", "contradicted_claims", "conflict", "partial",
            "best_cited_rank", "best_cited_relative_score", "unchecked_claims",
        }

    def test_there_is_no_percentage_in_the_output(self):
        assert set(assess(good())) == {"band", "reasons", "features"} and assess(good())["band"] in ("high", "medium", "low")


# ── Placeholders, failures and people ────────────────────


class TestPlaceholdersAndPeople:
    @pytest.mark.parametrize("state", [states.NO_EVIDENCE, states.NEEDS_SME, states.NEEDS_COMMERCIAL, states.NEEDS_LEGAL])
    def test_a_placeholder_is_low_because_a_person_has_to_supply_it(self, state):
        result = assess({"state": state, "evidence": []})
        assert result["band"] == "low" and codes(result) == ["no_answer"] and states.LABELS[state] in result["reasons"][0]["text"]

    def test_a_failure_is_low(self):
        assert codes(assess({"state": states.FAILED})) == ["failed"]

    def test_attach_adds_the_assessment_to_the_answer(self):
        attached = confidence.attach(good(), DOCS)
        assert attached["confidence"]["band"] == "high" and attached["state"] == states.FULLY


# ── In the drafting pipeline ─────────────────────────────


class TestPipeline:
    def test_every_answer_a_draft_produces_carries_a_band_and_reasons(self, agent, kb):  # noqa: F811
        kb.docs = {"data is encrypted at rest": [doc("Encryption", KB_A)]}
        agent.respond = lambda prompt: json.dumps(answer_of(requirement_id="R1"))
        requirements = [
            {"id": "R1", "text": "data is encrypted at rest"},
            {"id": "R2", "text": "Provide pricing for 500 users."},
            {"id": "R3", "text": "Are you ISO 9001 certified?"},
        ]

        results = list(strategies.stream_per_requirement(requirements=requirements, tone="professional", **ACCESS))

        assert [r["state"] for r in results] == ["fully_answered", "needs_commercial", "no_evidence"]
        assert all(r["confidence"]["band"] in ("high", "medium", "low") and r["confidence"]["reasons"] for r in results)
        assert [r["confidence"]["band"] for r in results[1:]] == ["low", "low"]
        assert results[0]["confidence"]["features"]["sources"] == 1

    def test_a_failed_generation_is_low_too(self, agent, kb, monkeypatch):  # noqa: F811
        kb.docs = {"data is encrypted at rest": [doc("Encryption", KB_A)]}
        agent.status = "failed"
        (result,) = strategies.stream_per_requirement(
            requirements=[{"id": "R1", "text": "data is encrypted at rest"}], tone="professional", **ACCESS)
        assert result["state"] == "failed" and result["confidence"]["band"] == "low"

    def test_the_span_carries_the_band_and_no_content(self):
        from app.telemetry import spans

        assert spans.safe_attributes({"confidence.band": "low", "confidence.reasons": "secret"}) == {"confidence.band": "low"}

# ── Stored with the answer, and calibrated ───────────────


def reply(requirement, band="low", state="partially_answered"):
    return {
        **answer_of(requirement_id=requirement["id"], answer_markdown=f"Answer for {requirement['id']}."),
        "state": state, "confidence": {"band": band, "reasons": [
            {"code": "partial", "effect": "lowers", "level": "medium", "text": "Only part of the requirement is covered"},
            {"code": "corroborated", "effect": "supports", "level": "high", "text": "2 independent sources cited"}],
            "features": {"sources": 2, "claims": 3, "nested": {"x": 1}, "long": "x" * 99}},
    }


class TestLedger:
    async def test_the_band_its_reasons_and_facts_are_stored_with_the_answer(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: reply(r, "low")
        project = await make_project(api)
        await draft(api, worker, project)

        first = (await state_of(api, project))["requirements"][0]

        assert first["confidence"] == "low"
        stored = first["answer"]["confidence"]
        assert stored["band"] == "low" and [r["code"] for r in stored["reasons"]] == ["partial", "corroborated"]
        assert stored["features"] == {"sources": 2, "claims": 3}  # nested and long values are not kept

    async def test_it_is_in_the_audit_trail_as_a_name_and_the_draft_status_counts_the_low_ones(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: reply(r, "low")
        script.hooks["R2"] = lambda r: reply(r, "high", "fully_answered")
        project = await make_project(api)
        await draft(api, worker, project)
        events = (await api.get(url(project, "events"), params={"rid": "R1"}, headers=A)).json()["events"]
        assert next(e for e in events if e["action"] == "answer.added")["details"]["confidence"] == "low"
        assert (await api.get(url(project, "draft"), headers=A)).json()["low_confidence"] == 1

    async def test_a_persons_text_has_no_band_and_cannot_be_given_one(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: reply(r, "low")
        project = await make_project(api)
        await draft(api, worker, project)

        written = await api.post(url(project, "requirements/R1/answers"), json={"answer_markdown": "Mine."}, headers=A)
        assert written.json()["confidence"] == "" and written.json()["answer"]["confidence"] == {}

        forged = await api.post(url(project, "requirements/R1/answers"),
                                json={"answer_markdown": "Mine.", "confidence": {"band": "high"}}, headers=A)
        assert forged.status_code == 422

    async def test_restoring_a_version_brings_its_band_back(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: reply(r, "low")
        project = await make_project(api)
        await draft(api, worker, project)
        await api.post(url(project, "requirements/R1/answers"), json={"answer_markdown": "Mine."}, headers=A)
        restored = await api.post(url(project, "requirements/R1/answers"), json={"restore_version": 1}, headers=A)
        assert restored.json()["confidence"] == "low"

    def test_what_is_kept_is_checked(self):
        assert service._confidence({"band": "great"}) == {} and service._confidence("high") == {} and service._confidence(None) == {}
        kept = service._confidence({"band": "high", "reasons": [
            {"code": "c", "effect": "lowers", "level": "low", "text": "t" * 999, "extra": "dropped"},
            {"effect": "sideways", "text": "x"}, {"effect": "lowers", "text": ""}, "str"], "features": {"a": [1]}})
        assert len(kept["reasons"]) == 1 and len(kept["reasons"][0]["text"]) == 200 and "extra" not in kept["reasons"][0]
        assert kept["features"] == {}


class TestCalibration:
    async def make(self, api, worker, script, bands):  # noqa: F811
        for rid, band in bands.items():
            script.hooks[rid] = lambda r, band=band: reply(r, band, "fully_answered")
        project = await make_project(api)
        await draft(api, worker, project)
        return project

    async def calibration(self, api, project, who=A):  # noqa: F811
        response = await api.get(url(project, "calibration"), headers=who)
        assert response.status_code == 200, response.text
        return response.json()

    async def test_each_band_is_scored_by_what_reviewers_did_with_its_answers(self, api, worker, script):  # noqa: F811
        project = await self.make(api, worker, script, {"R1": "high", "R2": "low", "R3": "low"})
        await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)  # accepted as drafted
        await api.post(url(project, "requirements/R2/answers"), json={"answer_markdown": "Rewritten."}, headers=A)  # edited

        result = await self.calibration(api, project)

        assert result["bands"]["high"] == {"answers": 1, "accepted": 1, "edited": 0, "pending": 0, "edit_rate": 0.0}
        assert result["bands"]["low"] == {"answers": 2, "accepted": 0, "edited": 1, "pending": 1, "edit_rate": 1.0}
        assert result["bands"]["medium"]["edit_rate"] is None and result["decided"] == 2

    async def test_an_answer_that_is_awaiting_review_is_pending_not_counted_either_way(self, api, worker, script):  # noqa: F811
        project = await self.make(api, worker, script, {"R1": "high"})
        result = await self.calibration(api, project)
        assert result["bands"]["high"]["pending"] == 1 and result["decided"] == 0

    async def test_placeholders_and_a_persons_own_answers_are_not_calibrated(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: reply(r, "low", "no_evidence")
        script.hooks["R3"] = lambda r: reply(r, "medium", "fully_answered")
        project = await make_project(api)
        await api.post(url(project, "requirements/R2/answers"), json={"answer_markdown": "Mine."}, headers=A)
        await draft(api, worker, project)
        assert sum(b["answers"] for b in (await self.calibration(api, project))["bands"].values()) == 1  # R3 only

    async def test_a_redraft_is_judged_by_its_own_outcome(self, api, worker, script):  # noqa: F811
        project = await self.make(api, worker, script, {"R1": "low"})
        await api.post(url(project, "requirements/R1/answers"), json={"answer_markdown": "Rewritten."}, headers=A)
        script.hooks["R1"] = lambda r: reply(r, "high", "fully_answered")
        await draft(api, worker, project, body={"rids": ["R1"]})
        bands = (await self.calibration(api, project))["bands"]
        assert bands["high"]["answers"] == 1 and bands["high"]["edited"] == 0 and bands["low"]["answers"] == 0

    async def test_restoring_the_models_own_version_is_not_an_edit(self, api, worker, script):  # noqa: F811
        project = await self.make(api, worker, script, {"R1": "high"})
        await api.post(url(project, "requirements/R1/answers"), json={"restore_version": 1}, headers=A)
        await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)
        assert (await self.calibration(api, project))["bands"]["high"]["accepted"] == 1

    async def test_only_people_who_can_open_the_project_can_see_it(self, api, worker, script):  # noqa: F811
        project = await self.make(api, worker, script, {"R1": "high"})
        assert (await api.get(url(project, "calibration"), headers=B)).status_code == 404


# ── Documents ────────────────────────────────────────────


RESPONSES = [
    {"requirement_id": "R1", "answer_markdown": "a", "meets": "Yes", "state": "fully_answered", "evidence": [],
     "confidence": assess(good())},
    {"requirement_id": "R2", "answer_markdown": "b", "meets": "Yes", "state": "partially_answered", "evidence": [],
     "confidence": assess(good(state=states.PARTIAL, verification=summary(1, unsupported=1)))},
    {"requirement_id": "R3", "answer_markdown": "c", "meets": "Yes", "evidence": []},
]


class TestDocuments:
    def test_the_matrix_has_a_confidence_column_last_with_the_reasons(self):
        sheet = load_workbook(io.BytesIO(generate_xlsx(RESPONSES, []))).active
        assert sheet.cell(1, 9).value == "Confidence" and sheet.cell(1, 8).value == "Verification"
        assert sheet.cell(2, 9).value.startswith("High: ")
        assert sheet.cell(3, 9).value.startswith("Low: 1 claim(s) are not supported") and "Only part of" in sheet.cell(3, 9).value
        assert sheet.cell(4, 9).value is None

    def test_the_document_gives_the_band_and_its_reasons(self):
        document = Document(io.BytesIO(generate_docx({"executive_summary": "s", "solution_overview": "o", "responses": RESPONSES})))
        text = [p.text for p in document.paragraphs]
        assert any(t.startswith("Confidence: High: ") for t in text)
        assert any(t.startswith("Confidence: Low: 1 claim(s) are not supported") for t in text)
        assert sum(1 for t in text if t.startswith("Confidence:")) == 2


# ── The page ─────────────────────────────────────────────


def with_band(band="low"):
    return requirement(
        status="drafted", state="partially_answered", confidence=band, answer_version=1,
        answer=answer(state="partially_answered", confidence=assess(good(state=states.PARTIAL, verification=summary(1, unsupported=1)))),
    )


class TestPage:
    def test_the_band_shows_in_the_header_and_the_answer_with_its_reasons_open_when_it_is_not_high(self):
        html = block(with_band())
        assert html.count("Low confidence") >= 2  # the header badge, and the one beside the answer
        assert "Why low confidence" in html and "<details style=\"margin-top:6px\" open>" in html
        assert "claim(s) are not supported by the evidence they cite" in html and "Only part of the requirement is covered" in html
        assert "&minus;</span>" in html and ">+</span>" in html  # what lowers it, and what supports it

    def test_a_high_band_keeps_its_reasons_closed(self):
        html = render("r.confidenceReasonsHtml(" + json.dumps(assess(good())) + ")")
        assert "Why high confidence" in html and " open" not in html.split("<summary")[0]

    def test_a_persons_answer_shows_no_band(self):
        html = block(requirement(status="edited", answer_version=1, answer=answer(source="user")))
        assert "confidence" not in html.lower()

    def test_an_unknown_band_shows_nothing_and_cannot_inject_a_style(self):
        assert render("r.confidenceBadgeHtml('x;background:red')") == "" and render("r.confidenceBadgeHtml(null)") == ""
        assert render("r.confidenceReasonsHtml({band: 'x', reasons: [{text: 't'}]})") == ""

    def test_reasons_are_escaped(self):
        hostile = '<img src=x onerror="alert(1)">'
        html = render("r.confidenceReasonsHtml({band: 'low', reasons: [{effect: 'lowers', text: " + json.dumps(hostile) + "}]})")
        assert hostile not in html and "&lt;img" in html

    def test_the_count_of_low_confidence_answers_is_shown(self):
        html = render("r.countsHtml({extracted: 1, low_confidence: 4}, 5)")
        assert "low confidence" in html and "<strong>4</strong>" in html


# ── On the golden set ────────────────────────────────────


@pytest.fixture(scope="module")
def runs():
    return {variant: run_golden(mode="replay", variant=variant) for variant in ("baseline", "degraded")}


class TestGoldenSet:
    def test_every_answer_on_the_golden_set_has_a_band_and_reasons(self, runs):
        answered = [s for r in runs.values() for s in r.scores if s.answered]
        assert answered and all(s.confidence in ("high", "medium", "low") and s.confidence_reasons is not None for s in answered)
        placeholders = [s for s in answered if s.state in states.PLACEHOLDER_STATES]
        assert placeholders and {s.confidence for s in placeholders} == {"low"}

    def test_the_low_band_is_wrong_more_often_than_the_high_band_on_answers_whose_truth_was_decided_in_advance(self, runs):
        # The injected set: 16 deliberately wrong claims and 16 right ones, labelled before any verifier saw them
        m = runs["baseline"].metrics
        assert m["injected_high_band_answers"] >= 10 and m["injected_low_band_answers"] >= 10
        assert m["injected_low_band_wrong_rate"] > m["injected_high_band_wrong_rate"]
        assert m["injected_high_band_wrong_rate"] == 0.0 and m["injected_confidence_separation"] >= 0.9
        assert {o.band for o in runs["baseline"].injections if o.expect == "right"} == {"high"}
        assert {o.band for o in runs["baseline"].injections if o.expect == "wrong"} == {"low"}

    def test_the_recorded_answers_of_the_golden_rfps_contain_no_wrong_answer_so_they_cannot_calibrate_the_band(self, runs):
        # Said plainly: every answer the model gave on the three RFPs passes the golden checks (the verifier's finer
        # flags are not labels, they are the band's input), so the edit rates here are all 0
        m = runs["baseline"].metrics
        assert m["confidence_high_answers"] >= 5 and m["confidence_low_answers"] >= 3
        assert m["confidence_high_edit_rate"] == 0.0 and m["confidence_low_edit_rate"] == 0.0

    def test_a_verifier_that_approves_everything_collapses_the_separation(self, tmp_path):
        import shutil

        golden = tmp_path / "golden"
        shutil.copytree(GOLDEN_DIR, golden)
        path = golden / "recordings" / "baseline" / "injections.json"
        recorded = json.loads(path.read_text(encoding="utf-8"))
        for record in recorded.values():
            record["text"] = json.dumps({"verdicts": [{"id": "C1", "verdict": "supported", "reason": "ok"}]})
        path.write_text(json.dumps(recorded), encoding="utf-8")

        lax = run_golden(mode="replay", variant="baseline", golden_dir=golden)

        # The rules still catch the claims they can without a model, so it does not fall to zero, but it falls well short
        assert lax.metrics["injected_confidence_separation"] < 0.7
        assert lax.metrics["injected_high_band_wrong_rate"] > 0.3

    def test_the_golden_clock_is_pinned_so_ages_do_not_drift(self, runs):
        from app.evalkit.run import AS_OF

        assert AS_OF == date(2026, 10, 1)
    def test_a_stale_historical_source_cited_for_a_claim_is_what_makes_a_golden_answer_low(self, runs):
        rt05 = next(s for s in runs["baseline"].scores if s.id == "RT-05")
        assert rt05.confidence == "low" and any("years old" in r for r in rt05.confidence_reasons)

    def test_the_report_shows_the_bands(self, runs):
        markdown = report.to_markdown(report.to_dict(runs["baseline"]))
        for label in ("Confidence separation", "High-band answers", "Low band: needs a change"):
            assert label in markdown

    def test_only_answers_count_never_placeholders(self):
        scores = [metrics.RequirementScore(case="c", id="x", text="t", supported=True, state=states.NO_EVIDENCE,
                                           confidence="low", needs_edit=True)]
        assert metrics.confidence_metrics(scores)["confidence_low_answers"] == 0

    def test_the_separation_is_unknown_until_both_bands_have_answers(self):
        one = [metrics.RequirementScore(case="c", id="x", text="t", supported=True, state=states.FULLY, confidence="high",
                                        needs_edit=False)]
        assert metrics.confidence_metrics(one)["confidence_separation"] is None

    def test_the_golden_check_for_14_weeks_accepts_14_week(self):
        with open("evals/golden/rfps/retail-xlsx.json", encoding="utf-8") as handle:
            spec = json.load(handle)
        rt05 = next(r for r in spec["requirements"] if r["id"] == "RT-05")
        assert ["14 weeks", "14-week"] in rt05["contains_any"]

# ── Found in review ──────────────────────────────────────


class TestUncheckedIsNotUnsupported:
    def unchecked(self, claims=2):
        return {"status": "unchecked", "claims": claims, "supported": 0, "partial": 0, "unsupported": 0, "contradicted": 0,
                "unchecked": claims, "by_rule": 0}

    def test_claims_nobody_judged_cap_the_band_at_medium_not_low(self):
        result = assess(good(verification=self.unchecked()))
        assert result["band"] == "medium" and codes(result) == ["unchecked"]
        assert result["features"]["supported_fraction"] is None and result["features"]["unchecked_claims"] == 2

    def test_a_web_only_claim_beside_a_supported_one_does_not_make_the_answer_low(self):
        mixed = {"status": "partially_verified", "claims": 2, "supported": 1, "partial": 0, "unsupported": 0,
                 "contradicted": 0, "unchecked": 1, "by_rule": 0}
        result = assess(good(verification=mixed))
        assert result["band"] == "medium" and codes(result) == ["partly_unchecked"]
        assert result["features"]["supported_fraction"] == 1.0  # of the claims that were judged

    def test_the_share_is_still_taken_over_what_was_judged(self):
        judged = {"status": "partially_verified", "claims": 4, "supported": 1, "partial": 1, "unsupported": 0,
                  "contradicted": 0, "unchecked": 2, "by_rule": 0}
        assert "few_claims_supported" in codes(assess(good(verification=judged)))  # 1 of 2 judged is 50%


class TestCalibrationOfSetAsideAnswers:
    async def test_bringing_back_an_older_model_version_sets_the_latest_one_aside(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: reply(r, "low", "fully_answered")
        project = await make_project(api)
        await draft(api, worker, project)
        script.hooks["R1"] = lambda r: reply(r, "high", "fully_answered")
        await draft(api, worker, project, body={"rids": ["R1"]})  # v2: High
        await api.post(url(project, "requirements/R1/answers"), json={"restore_version": 1}, headers=A)  # v3
        await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)

        bands = (await api.get(url(project, "calibration"), headers=A)).json()["bands"]

        assert bands["high"]["edited"] == 1 and bands["high"]["pending"] == 0 and bands["high"]["edit_rate"] == 1.0


class TestTheEvaluationLabelIsNotMadeOfTheBandsInputs:
    def test_an_invented_figure_alone_does_not_make_an_answer_need_a_change(self):
        score = metrics.score_requirement(
            case="c", expected={"id": "x", "text": "t", "supported": True}, extracted=True, docs=[],
            answer={"answer_markdown": "Fine.", "state": "fully_answered", "meets": "Yes", "evidence": []},
            raw_text="We cut costs by 40%.", sources="", retrieval_rank=None, latency_s=None, usage=None,
        )
        assert score.unsupported_raw == ["40%"] and score.needs_edit is False
