"""Claims and their verification in the ledger, the approval block, the documents and the page (part 2 of #35)."""

import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook

from app.agents import states, verify
from app.ledger import drafting, service
from app.ledger import store as ledger_store
from app.outputs.docx_writer import generate_docx
from app.outputs.xlsx_writer import generate_xlsx

from .test_answer_states import answer_of
from .test_authorization_api import A
from .test_draft_jobs import worker  # noqa: F401  (a fixture)
from .test_ledger_api import api, api_settings, make_project, url  # noqa: F401  (api is a fixture)
from .test_ledger_drafting import draft, script, state_of  # noqa: F401  (script is a fixture)
from .test_ui_projects import answer, block, render, requirement

EVIDENCE = [
    {"id": "E1", "title": "Policy", "source_url": "/kb/doc/abc", "quote": "AES-256 at rest", "quote_verified": True},
    {"id": "E2", "title": "Old", "source_url": "/kb/doc/def", "quote": "made up", "quote_verified": False},
]
CLAIMS = [
    {"id": "C1", "text": "Data is encrypted with AES-256.", "evidence_ids": ["E1"], "verdict": "supported",
     "reason": "Stated.", "checked_by": "model"},
    {"id": "C2", "text": "Keys are rotated daily.", "evidence_ids": ["E2"], "verdict": "unsupported",
     "reason": "The quoted text is not in the cited source.", "checked_by": "rule"},
]
FLAGGED = {"status": "flagged", "claims": 2, "supported": 1, "partial": 0, "unsupported": 1, "contradicted": 0,
           "unchecked": 0, "by_rule": 1, "model": "gpt-4o"}
VERIFIED = {"status": "verified", "claims": 1, "supported": 1, "partial": 0, "unsupported": 0, "contradicted": 0,
            "unchecked": 0, "by_rule": 0, "model": "gpt-4o"}


def verified_reply(requirement, *, flagged=False):
    return {
        **answer_of(requirement_id=requirement["id"], answer_markdown=f"Answer for {requirement['id']}."),
        "evidence": EVIDENCE if flagged else EVIDENCE[:1],
        "claims": CLAIMS if flagged else CLAIMS[:1],
        "verification": FLAGGED if flagged else VERIFIED,
        "state": "partially_answered" if flagged else "fully_answered",
    }


# ── Stored with the answer ───────────────────────────────


class TestStoredInTheLedger:
    async def test_every_claim_keeps_its_evidence_ids_and_verdict_with_the_answer(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r, flagged=True)
        project = await make_project(api)
        await draft(api, worker, project)

        first = (await state_of(api, project))["requirements"][0]

        assert first["verification"] == "flagged"
        stored = first["answer"]
        assert [(c["id"], c["evidence_ids"], c["verdict"], c["checked_by"]) for c in stored["claims"]] == [
            ("C1", ["E1"], "supported", "model"), ("C2", ["E2"], "unsupported", "rule"),
        ]
        assert stored["verification"]["status"] == "flagged" and stored["verification"]["unsupported"] == 1
        assert [(e["id"], e["quote_verified"]) for e in stored["evidence"]] == [("E1", True), ("E2", False)]

    async def test_the_claims_are_in_the_version_history_too(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r, flagged=True)
        project = await make_project(api)
        await draft(api, worker, project)

        (version,) = (await api.get(url(project, "requirements/R1/answers"), headers=A)).json()["answers"]
        assert len(version["claims"]) == 2 and version["verification"]["status"] == "flagged"

    async def test_the_audit_trail_records_the_status_but_never_the_claims(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r, flagged=True)
        project = await make_project(api)
        await draft(api, worker, project)

        events = (await api.get(url(project, "events"), params={"rid": "R1"}, headers=A)).json()["events"]
        added = next(e for e in events if e["action"] == "answer.added")
        assert added["details"]["verification"] == "flagged"
        assert "keys are rotated" not in json.dumps(events).lower()

    async def test_an_answer_without_claims_has_none_and_no_verification(self, api, worker, script):  # noqa: F811
        project = await make_project(api)
        await draft(api, worker, project)
        first = (await state_of(api, project))["requirements"][0]
        assert first["verification"] == "" and first["answer"]["claims"] == [] and first["answer"]["verification"] == {}

    def test_what_is_kept_is_cleaned_and_bounded(self):
        claims = service._claims(
            [{"text": "x" * 5000, "evidence_ids": ["E1", 5, "E2"] + ["E9"] * 20, "verdict": "great", "reason": "r" * 999,
              "checked_by": "magic"}, {"text": ""}, "str", None] + [{"text": "n"}] * 40
        )
        assert len(claims) == 22 and len(service._claims([{"text": "n"}] * 40)) == verify.MAX_CLAIMS
        assert len(claims[0]["text"]) == verify.MAX_CLAIM_CHARS and len(claims[0]["evidence_ids"]) == 9  # ten read, one of them not a string
        assert claims[0]["verdict"] is None and claims[0]["checked_by"] == "" and len(claims[0]["reason"]) == 200
        assert service._verification({"status": "bogus"}) == {} and service._verification("flagged") == {}
        kept = service._verification({**FLAGGED, "supported": -3, "claims": True, "model": "m" * 999, "extra": "dropped"})
        assert kept["supported"] == 0 and kept["claims"] == 0 and len(kept["model"]) == 200 and "extra" not in kept

    def test_the_evidence_id_and_the_quote_check_survive_but_nothing_else_is_trusted(self):
        (kept,) = service._evidence([{"id": "E1" * 20, "quote_verified": "yes", "title": "t", "source_url": "u", "quote": "q"}])
        assert kept["id"] == ("E1" * 20)[:10] and "quote_verified" not in kept


# ── Approval ─────────────────────────────────────────────


class TestApproval:
    async def test_an_answer_with_a_claim_its_evidence_does_not_support_cannot_be_approved(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r, flagged=True)
        project = await make_project(api)
        await draft(api, worker, project)

        refused = await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)

        assert refused.status_code == 409 and "1 claim(s)" in refused.json()["detail"]
        assert (await api.get(url(project, "requirements/R1"), headers=A)).json()["status"] == "drafted"

    async def test_a_verified_answer_can_be_approved(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r)
        project = await make_project(api)
        await draft(api, worker, project)
        approved = await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)
        assert approved.status_code == 200 and approved.json()["verification"] == "verified"

    async def test_a_person_fixing_the_answer_lifts_the_block_and_owns_the_text(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r, flagged=True)
        project = await make_project(api)
        await draft(api, worker, project)

        fixed = await api.post(url(project, "requirements/R1/answers"), json={"answer_markdown": "Keys rotate every 90 days."}, headers=A)

        body = fixed.json()
        assert body["verification"] == "" and body["answer"]["claims"] == [] and body["answer"]["verification"] == {}
        assert (await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)).status_code == 200

    async def test_restoring_the_flagged_version_brings_the_block_back(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r, flagged=True)
        project = await make_project(api)
        await draft(api, worker, project)
        await api.post(url(project, "requirements/R1/answers"), json={"answer_markdown": "Fixed."}, headers=A)

        restored = await api.post(url(project, "requirements/R1/answers"), json={"restore_version": 1}, headers=A)

        assert restored.json()["verification"] == "flagged"
        assert (await api.patch(url(project, "requirements/R1"), json={"status": "approved"}, headers=A)).status_code == 409

    async def test_a_user_cannot_attach_claims_to_their_own_text(self, api):  # noqa: F811
        project = await make_project(api)
        response = await api.post(
            url(project, "requirements/R1/answers"),
            json={"answer_markdown": "Mine.", "claims": CLAIMS, "verification": VERIFIED}, headers=A,
        )
        assert response.status_code == 422  # unknown fields are refused

    async def test_the_draft_status_counts_the_flagged_requirements(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r, flagged=True)
        script.hooks["R2"] = lambda r: verified_reply(r)
        project = await make_project(api)
        await draft(api, worker, project)
        status = (await api.get(url(project, "draft"), headers=A)).json()
        assert status["flagged"] == 1

    async def test_the_documents_are_built_with_each_answers_verification(self, api, worker, script):  # noqa: F811
        script.hooks["R1"] = lambda r: verified_reply(r, flagged=True)
        project = await make_project(api)
        await draft(api, worker, project)
        sink = drafting.ProjectSink(
            ledger_store.store_for(api_settings()), api_settings(), project["project"]["id"], "j",
            drafting.select(ledger_store.store_for(api_settings()), project["project"]["id"], ["R1"]),
        )
        responses, _ = sink.document_inputs()
        by_id = {r["requirement_id"]: r for r in responses}
        assert by_id["R1"]["verification"]["status"] == "flagged" and len(by_id["R1"]["claims"]) == 2


# ── Documents ────────────────────────────────────────────


RESPONSES = [
    {"requirement_id": "R1", "answer_markdown": "a", "meets": "Yes", "state": "fully_answered", "evidence": [],
     "verification": VERIFIED, "claims": CLAIMS[:1]},
    {"requirement_id": "R2", "answer_markdown": "b", "meets": "Yes", "state": "partially_answered", "evidence": [],
     "verification": FLAGGED, "claims": CLAIMS},
    {"requirement_id": "R3", "answer_markdown": "c", "meets": "Yes", "evidence": []},
]


class TestDocuments:
    def test_the_matrix_has_a_verification_column_after_the_state_and_lists_flagged_claims_in_the_notes(self):
        sheet = load_workbook(io.BytesIO(generate_xlsx(RESPONSES, []))).active
        assert [sheet.cell(1, c).value for c in (7, 8)] == ["Answer state", "Verification"]
        assert [sheet.cell(r, 8).value for r in (2, 3, 4)] == [
            "Verified (1 of 1 claims supported)", "Flagged (1 unsupported)", None]
        assert "Unsupported claim: Keys are rotated daily." in sheet.cell(3, 5).value
        assert "Keys are rotated" not in (sheet.cell(2, 5).value or "")

    def test_the_document_says_whether_each_answer_was_verified(self):
        document = Document(io.BytesIO(generate_docx({"executive_summary": "s", "solution_overview": "o", "responses": RESPONSES})))
        text = [p.text for p in document.paragraphs]
        assert "Verification: Verified (1 of 1 claims supported)" in text
        assert "Verification: Flagged (1 unsupported)" in text
        assert sum(1 for t in text if t.startswith("Verification:")) == 2  # none for an answer that was never checked

    @pytest.mark.parametrize("summary, words", [
        ({"status": "verified", "claims": 3, "supported": 3}, "Verified (3 of 3 claims supported)"),
        ({"status": "partially_verified", "claims": 3, "supported": 2, "partial": 1}, "Partly verified (2 of 3 claims supported)"),
        ({"status": "flagged", "unsupported": 1, "contradicted": 2}, "Flagged (1 unsupported, 2 contradicted)"),
        ({"status": "unchecked", "reason": "x"}, "Not checked"),
        ({"status": "other"}, ""), (None, ""), ("flagged", ""),
    ])
    def test_a_summary_in_words(self, summary, words):
        assert verify.describe(summary) == words


# ── The page ─────────────────────────────────────────────


def flagged_requirement():
    return requirement(
        status="drafted", state="partially_answered", verification="flagged", answer_version=1,
        answer=answer(state="partially_answered", claims=CLAIMS, evidence=EVIDENCE, verification=FLAGGED),
    )


class TestPage:
    def test_a_flagged_answer_shows_its_badge_the_unsupported_claims_open_and_the_evidence_it_could_not_find(self):
        html = block(flagged_requirement())
        assert "Flagged (1 unsupported)" in html
        assert "<details" in html and " open>" in html and "Claims (2)" in html
        assert "Keys are rotated daily." in html and "The quoted text is not in the cited source." in html
        assert "(checked without a model)" in html and "quote not found in the source" in html
        assert "<strong style=\"color:var(--text-muted)\">E1</strong>" in html

    def test_a_flagged_answer_cannot_be_approved_on_the_page_either(self):
        assert 'data-action="approve"' not in block(flagged_requirement())
        assert 'data-action="approve"' in block(requirement(
            status="drafted", state="fully_answered", verification="verified", answer_version=1,
            answer=answer(claims=CLAIMS[:1], evidence=EVIDENCE[:1], verification=VERIFIED)))

    def test_a_verified_answer_keeps_its_claims_closed(self):
        html = block(requirement(status="drafted", state="fully_answered", verification="verified", answer_version=1,
                                 answer=answer(claims=CLAIMS[:1], evidence=EVIDENCE[:1], verification=VERIFIED)))
        assert "Verified (1 of 1" not in html and ">Verified<" in html  # the badge says Verified
        assert "<details" in html and " open>" not in html

    def test_an_answer_that_was_not_checked_says_so_and_a_persons_answer_shows_nothing(self):
        unchecked = {"status": "unchecked", "reason": "The verifier could not be run."}
        html = block(requirement(status="drafted", answer_version=1, answer=answer(verification=unchecked)))
        assert "Not checked" in html and "The verifier could not be run." in html
        assert "Claims (" not in block(requirement(status="edited", answer_version=1, answer=answer(source="user")))

    def test_claim_text_and_reasons_are_escaped(self):
        hostile = '<img src=x onerror="alert(1)">'
        html = block(requirement(status="drafted", answer_version=1, answer=answer(
            claims=[{"id": "C1", "text": hostile, "evidence_ids": [hostile], "verdict": "supported", "reason": hostile}],
            verification={**VERIFIED, "reason": hostile})))
        assert hostile not in html and "&lt;img" in html

    def test_an_unknown_verdict_or_status_cannot_inject_a_style(self):
        html = render("r.verificationBadgeHtml({status: 'x;background:red'})")
        assert html == ""
        html = block(requirement(status="drafted", answer_version=1, answer=answer(
            claims=[{"text": "t", "verdict": "x;background:red", "evidence_ids": []}])))
        assert "x;background:red</strong>" in html and 'style="x' not in html and "style=\"color:var(--text-muted)\">x;" in html

    def test_counts_include_the_flagged_requirements(self):
        html = render("r.countsHtml({extracted: 1, flagged: 2}, 3)")
        assert "flagged" in html and "<strong>2</strong>" in html

    def test_a_claim_with_no_evidence_is_called_out(self):
        html = block(requirement(status="drafted", answer_version=1, answer=answer(
            claims=[{"id": "C1", "text": "We hold HITRUST.", "evidence_ids": [], "verdict": "unsupported", "reason": "No evidence."}])))
        assert "[no evidence]" in html


def test_the_states_and_verification_statuses_are_distinct_vocabularies():
    assert not set(verify.STATUSES) & set(states.STATES)
