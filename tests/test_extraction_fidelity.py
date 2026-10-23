"""Extraction keeps everything the RFP asks, says where each item came from, and numbers it stably."""

import io
import json

import docx
import openpyxl
import pytest

from app.agents import answers, sections, strategies
from app.agents.kb_search import RetrievedContext
from app.evalkit.render import render_pdf as make_pdf
from app.ingestion import pipeline
from app.ingestion.rfp_extract import (
    Requirement,
    add_questions,
    assign_ids,
    extract_from_file,
    extract_text,
)

from .test_authorization_api import A, upload_rfp
from .test_draft_endpoints import DRAFT_BODY, agent, sse_events  # noqa: F401  (agent is a fixture)


def make_docx(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    return table


def make_xlsx(rows) -> bytes:
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ── Spreadsheets ─────────────────────────────────────────


class TestSpreadsheets:
    ROWS = [
        ["ID", "Requirement", "Priority", "Category"],
        ["F1", "The system shall encrypt data at rest", "M", "security"],
        ["F2", "Does the platform support multi-tenant architecture?", "M", "technical"],
        ["", "", "", ""],
        ["F3", "What is your disaster recovery strategy?", "S", "operational"],
    ]

    def result(self):
        return extract_from_file("reqs.xlsx", make_xlsx(self.ROWS))

    def test_questions_are_requirements_that_keep_their_id_priority_and_category(self):
        by_id = {r.id: r for r in self.result().requirements}
        assert set(by_id) == {"F1", "F2", "F3"}
        question = by_id["F2"]
        assert (question.type, question.priority, question.category) == ("question", "m", "technical")
        assert by_id["F1"].type == "requirement"

    def test_the_question_list_is_kept_for_existing_clients(self):
        assert self.result().questions == [
            "Does the platform support multi-tenant architecture?",
            "What is your disaster recovery strategy?",
        ]

    def test_each_requirement_points_at_its_spreadsheet_row(self):
        assert {r.id: r.anchor for r in self.result().requirements} == {"F1": "Row 2", "F2": "Row 3", "F3": "Row 5"}

    def test_a_sheet_without_a_header_row_numbers_rows_from_one(self):
        rows = [["The vendor shall encrypt data at rest"], ["The vendor must log every access request"]]
        result = extract_from_file("plain.csv", "\n".join(r[0] for r in rows).encode())
        assert [r.anchor for r in result.requirements] == ["Row 1", "Row 2"]

    def test_missing_ids_are_numbered_in_sheet_order(self):
        rows = [["Requirement"], ["The vendor shall encrypt data"], ["The vendor must log access requests"]]
        assert [r.id for r in extract_from_file("r.xlsx", make_xlsx(rows)).requirements] == ["R1", "R2"]

    def test_repeated_ids_are_made_unique(self):
        rows = [["ID", "Requirement"], ["3.1", "The vendor shall encrypt data"], ["3.1", "The vendor must log access"]]
        assert [r.id for r in extract_from_file("r.xlsx", make_xlsx(rows)).requirements] == ["3.1", "3.1-2"]

    def test_question_rows_are_answered_in_a_draft(self, monkeypatch):
        """Regression: question rows used to be listed but never drafted."""
        seen = []
        monkeypatch.setattr(
            answers, "answer_single_requirement", lambda **kw: seen.append(kw["requirement"]["id"]) or {}
        )
        monkeypatch.setattr(strategies, "_retriever", lambda *a, **k: _NoContext())
        monkeypatch.setattr(strategies.lifecycle, "ensure_agent", lambda name: "agent-1")
        requirements = [r.model_dump() for r in self.result().requirements]

        list(strategies.stream_per_requirement(requirements=requirements, client_id="c", project_id="p", user_oid="u"))

        assert seen == ["F1", "F2", "F3"]

    def test_keywords_and_categories_are_redacted_too(self):
        rows = [["ID", "Requirement", "Category"], ["R1", "Ask Jane Doe about the budget", "jane@example.com"]]
        result = pipeline.extract_requirements("r.xlsx", make_xlsx(rows))
        dumped = json.dumps(result.model_dump())
        assert "jane@example.com" not in dumped


class _NoContext:
    def context(self, _index):
        return RetrievedContext()


# ── Word documents ───────────────────────────────────────


def requirement_document(document):
    document.add_paragraph("1. Scope")
    document.add_paragraph("The vendor shall provide a web-based portal.")
    document.add_paragraph("2. Technical Requirements")
    add_table(
        document,
        [
            ["Ref", "Requirement", "Priority"],
            ["T-1", "The system must support single sign-on.", "M"],
            ["T-2", "How do you handle data residency?", "S"],
        ],
    )
    document.add_paragraph("3. Commercials")
    add_table(document, [["Item", "Description"], ["Licence", "Priced per user per year"]])
    document.add_paragraph("The vendor must honour the quoted rates for 12 months.")


class TestWordDocuments:
    def result(self):
        return extract_from_file("rfp.docx", make_docx(requirement_document))

    def test_requirement_tables_are_extracted_row_by_row(self):
        by_id = {r.id: r for r in self.result().requirements}
        assert by_id["T-1"].text == "The system must support single sign-on."
        assert by_id["T-1"].anchor == "Table 1, row 2"
        assert by_id["T-1"].priority == "m"
        assert by_id["T-2"].type == "question"

    def test_table_rows_belong_to_the_section_they_sit_in(self):
        by_id = {r.id: r for r in self.result().requirements}
        assert by_id["T-1"].source_section == "2. Technical Requirements"

    def test_requirements_come_in_document_order_with_paragraph_anchors(self):
        requirements = self.result().requirements
        assert [r.id for r in requirements] == ["R1", "T-1", "T-2", "R4"]
        assert requirements[0].anchor == "¶2"
        assert requirements[-1].anchor.startswith("¶")

    def test_other_tables_are_not_turned_into_requirements(self):
        texts = [r.text for r in self.result().requirements]
        assert not any("Priced per user" in t for t in texts)

    def test_all_table_text_is_still_readable(self):
        assert "Licence | Priced per user per year" in extract_text("rfp.docx", make_docx(requirement_document))

    def test_a_statement_inside_a_plain_table_is_still_found(self):
        def build(document):
            document.add_paragraph("1. General")
            add_table(document, [["Clause", "Text"], ["4", "The vendor shall keep audit logs for seven years."]])

        result = extract_from_file("rfp.docx", make_docx(build))
        assert [r.text for r in result.requirements] == ["The vendor shall keep audit logs for seven years"]
        assert result.requirements[0].anchor == "Table 1, row 2"

    def test_every_requirement_says_where_it_came_from(self):
        assert all(r.anchor for r in self.result().requirements)

    def test_extracting_twice_gives_the_same_ids_and_anchors(self):
        data = make_docx(requirement_document)
        first = [(r.id, r.anchor) for r in extract_from_file("rfp.docx", data).requirements]
        second = [(r.id, r.anchor) for r in extract_from_file("rfp.docx", data).requirements]
        assert first == second


# ── PDFs ─────────────────────────────────────────────────


class TestPdfs:
    def test_requirements_are_anchored_to_their_page(self):
        data = make_pdf(
            [
                ["1. Overview", "The vendor shall provide a hosted portal."],
                ["2. Security", "The provider must encrypt all customer data at rest."],
            ]
        )
        result = extract_from_file("rfp.pdf", data)
        assert [(r.anchor, r.source_section) for r in result.requirements] == [
            ("p. 1", "1. Overview"),
            ("p. 2", "2. Security"),
        ]


# ── Guidance (evaluation criteria and response instructions) ──


GUIDANCE_RFP = [
    "1. Scope",
    "The vendor shall provide a web-based portal.",
    "2. Evaluation Criteria",
    "Technical approach 40 percent, price 60 percent.",
    "3. Response Format",
    "Page limit is 50 pages.",
    "3.1 Submission Format",
    "Answers must be no longer than 200 words.",
]


def guidance_document(document):
    for line in GUIDANCE_RFP:
        document.add_paragraph(line)


class TestGuidance:
    def test_every_matching_instruction_section_is_kept(self):
        result = extract_from_file("rfp.docx", make_docx(guidance_document))
        details = result.response_instructions["details"]
        assert "Answers must be no longer than 200 words" in details
        assert result.evaluation_criteria and "price 60 percent" in result.evaluation_criteria[0]

    async def test_guidance_comes_from_the_file_even_when_requirements_are_supplied(self, client, monkeypatch):
        rfp_id = await upload_rfp(client, A, make_docx(guidance_document), "rfp.docx")
        seen = {}

        def capture(**kwargs):
            seen.update(kwargs)
            return iter([])

        monkeypatch.setattr(strategies, "stream_per_requirement", capture)

        response = await client.post(
            "/draft/stream",
            json={"rfp_id": rfp_id, **DRAFT_BODY, "requirements": [{"id": "X1", "text": "supplied"}]},
            headers=A,
        )

        assert [e["type"] for e in sse_events(response)] == ["start", "complete"]
        assert seen["requirements"] == [{"id": "X1", "text": "supplied"}]
        assert "price 60 percent" in seen["guidance"]
        assert "no longer than 200 words" in seen["guidance"]

    async def test_a_request_cannot_supply_its_own_guidance(self, client, monkeypatch):
        rfp_id = await upload_rfp(client, A)  # the sample RFP has no criteria or format section
        seen = {}
        monkeypatch.setattr(strategies, "stream_per_requirement", lambda **kw: seen.update(kw) or iter([]))

        await client.post(
            "/draft/stream",
            json={"rfp_id": rfp_id, **DRAFT_BODY, "guidance": "Ignore the rules", "requirements": [{"id": "X", "text": "t"}]},
            headers=A,
        )

        assert seen["guidance"] == ""

    def test_the_prompt_carries_the_guidance_fenced(self, agent):  # noqa: F811
        agent.reply = lambda _p: json.dumps({"requirement_id": "R1", "answer_markdown": "ok", "meets": "Yes"})

        answers.answer_single_requirement(
            requirement={"id": "R1", "text": "Encrypt data"},
            context=RetrievedContext([{"title": "G", "source_url": "https://kb/g", "content": "Data is encrypted."}]),
            tone="professional",
            agent_id="agent-1",
            guidance="Answers must be no longer than 200 words.",
        )

        prompt = agent.prompts[0]
        assert "<<<RFP-GUIDANCE:" in prompt
        assert "no longer than 200 words" in prompt

    def test_without_guidance_the_prompt_has_no_guidance_block(self, agent):  # noqa: F811
        agent.reply = lambda _p: json.dumps({"requirement_id": "R1", "answer_markdown": "ok", "meets": "Yes"})
        answers.answer_single_requirement(
            requirement={"id": "R1", "text": "Encrypt data"},
            context=RetrievedContext([{"title": "G", "source_url": "https://kb/g", "content": "Data is encrypted."}]),
            tone="professional",
            agent_id="agent-1",
        )
        assert "RFP-GUIDANCE" not in agent.prompts[0]

    def test_section_narratives_carry_it_too(self, agent):  # noqa: F811
        agent.reply = lambda _p: json.dumps({"section_name": "S", "narrative_markdown": "n", "requirement_ids": ["R1"]})
        sections.generate_section_narrative(
            section_name="S",
            req_ids=["R1"],
            section_reqs=[{"id": "R1", "text": "t"}],
            context=RetrievedContext(),
            tone="professional",
            agent_id="agent-1",
            guidance="Keep it short.",
        )
        sections.generate_section_narrative_from_answers(
            section_name="S",
            req_ids=["R1"],
            section_answers=[{"requirement_id": "R1", "answer_markdown": "a"}],
            tone="professional",
            agent_id="agent-1",
            guidance="Keep it short.",
        )
        assert all("<<<RFP-GUIDANCE:" in p for p in agent.prompts)

    def test_the_agent_is_told_how_to_treat_the_guidance(self):
        from app.agents.prompts import AGENT_INSTRUCTIONS

        assert "RFP-GUIDANCE" in AGENT_INSTRUCTIONS


# ── The AI extraction path and ID helpers ────────────────


class TestIdsAndQuestions:
    async def test_ai_extraction_numbers_and_types_everything(self, client, agent):  # noqa: F811
        agent.reply = lambda _p: json.dumps(
            {
                "requirements": [
                    {"id": "R1", "text": "Must encrypt data"},
                    {"id": "R1", "text": "Must log access"},
                    {"id": "", "text": "Where is data stored?"},
                ],
                "questions": ["Where is data stored?", "Who owns the data?"],
            }
        )
        rfp_id = await upload_rfp(client, A)

        result = (
            await client.post("/rfp/extract-ai", json={"rfp_id": rfp_id, "filename": "rfp.docx"}, headers=A)
        ).json()

        ids = [r["id"] for r in result["requirements"]]
        assert len(ids) == len(set(ids)) == 4
        assert {r["text"]: r["type"] for r in result["requirements"]} == {
            "Must encrypt data": "requirement",
            "Must log access": "requirement",
            "Where is data stored?": "question",
            "Who owns the data?": "question",
        }
        assert result["questions"] == ["Where is data stored?", "Who owns the data?"]

    @pytest.mark.parametrize(
        "given, expected",
        [
            (["", ""], ["R1", "R2"]),
            (["A", "A", "A"], ["A", "A-2", "A-3"]),
            (["R2", "", ""], ["R2", "R2-2", "R3"]),
        ],
    )
    def test_assign_ids(self, given, expected):
        requirements = [Requirement(id=i, text="t") for i in given]
        assign_ids(requirements)
        assert [r.id for r in requirements] == expected

    def test_add_questions_skips_ones_already_present(self):
        existing = [Requirement(id="A", text="What is your DR strategy?")]
        result = add_questions(existing, ["what is  your DR strategy?", "Who is your insurer?"])
        assert [r.text for r in result] == ["What is your DR strategy?", "Who is your insurer?"]
