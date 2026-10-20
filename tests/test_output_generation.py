"""Tests for DOCX and XLSX output generation."""

import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook

from app.outputs.docx_writer import generate_docx
from app.outputs.xlsx_writer import generate_xlsx


# ── Fixtures ─────────────────────────────────────────────

SAMPLE_DRAFT = {
    "executive_summary": "We propose a comprehensive cloud solution.",
    "solution_overview": "Our platform is built on Azure with high availability.",
    "responses": [
        {
            "requirement_id": "REQ-001",
            "answer_markdown": "Our platform provides 99.9% uptime backed by Azure SLA.",
            "meets": "Yes",
            "evidence": [
                {
                    "title": "Azure SLA Documentation",
                    "source_url": "https://azure.microsoft.com/sla/",
                    "quote": "Azure guarantees 99.9% uptime for compute services.",
                }
            ],
            "assumptions": ["Customer has Azure subscription"],
            "risks": [],
            "followups": ["Confirm SLA tier required"],
        },
        {
            "requirement_id": "REQ-002",
            "answer_markdown": "RBAC is fully supported via Azure AD integration.",
            "meets": "Yes",
            "evidence": [
                {
                    "title": "RBAC Guide",
                    "source_url": "https://docs.example.com/rbac",
                    "quote": "",
                }
            ],
            "assumptions": [],
            "risks": ["Custom roles may require additional configuration"],
            "followups": [],
        },
        {
            "requirement_id": "REQ-003",
            "answer_markdown": "Not provided in sources – recommend human review.",
            "meets": "No",
            "evidence": [],
            "assumptions": [],
            "risks": [],
            "followups": ["Clarify specific integration requirements"],
        },
    ],
}

SAMPLE_REQUIREMENTS = [
    {"id": "REQ-001", "text": "The vendor shall ensure 99.9% uptime.", "priority": "mandatory"},
    {"id": "REQ-002", "text": "The vendor shall implement RBAC.", "priority": "mandatory"},
    {"id": "REQ-003", "text": "The vendor must integrate with legacy systems.", "priority": "standard"},
]


# ── DOCX tests ───────────────────────────────────────────


class TestDocxGeneration:
    def test_generates_valid_docx_bytes(self):
        data = generate_docx(SAMPLE_DRAFT)
        assert isinstance(data, bytes)
        assert len(data) > 100

    def test_docx_is_openable(self):
        data = generate_docx(SAMPLE_DRAFT)
        doc = Document(io.BytesIO(data))
        assert len(doc.paragraphs) > 0

    def test_docx_contains_executive_summary(self):
        data = generate_docx(SAMPLE_DRAFT)
        doc = Document(io.BytesIO(data))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "comprehensive cloud solution" in full_text

    def test_docx_contains_requirement_headings(self):
        data = generate_docx(SAMPLE_DRAFT)
        doc = Document(io.BytesIO(data))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "REQ-001" in full_text
        assert "REQ-002" in full_text
        assert "REQ-003" in full_text

    def test_docx_has_compliance_status(self):
        data = generate_docx(SAMPLE_DRAFT)
        doc = Document(io.BytesIO(data))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Compliance: Yes" in full_text
        assert "Compliance: No" in full_text

    def test_docx_handles_empty_draft(self):
        data = generate_docx({"responses": []})
        assert isinstance(data, bytes)
        doc = Document(io.BytesIO(data))
        assert len(doc.paragraphs) > 0


# ── XLSX tests ───────────────────────────────────────────


class TestXlsxGeneration:
    def test_generates_valid_xlsx_bytes(self):
        data = generate_xlsx(SAMPLE_DRAFT["responses"], SAMPLE_REQUIREMENTS)
        assert isinstance(data, bytes)
        assert len(data) > 100

    def test_xlsx_has_correct_headers(self):
        data = generate_xlsx(SAMPLE_DRAFT["responses"], SAMPLE_REQUIREMENTS)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, 7)]
        assert "Requirement ID" in headers
        assert "Requirement Text" in headers
        assert "Priority" in headers
        assert "Meets (Yes/Partial/No)" in headers
        assert "Evidence / Notes" in headers
        assert "Source URL" in headers

    def test_xlsx_has_correct_row_count(self):
        data = generate_xlsx(SAMPLE_DRAFT["responses"], SAMPLE_REQUIREMENTS)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        # 1 header + 3 data rows
        assert ws.max_row == 4

    def test_xlsx_requirement_ids_match(self):
        data = generate_xlsx(SAMPLE_DRAFT["responses"], SAMPLE_REQUIREMENTS)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        req_ids = [ws.cell(row=r, column=1).value for r in range(2, 5)]
        assert "REQ-001" in req_ids
        assert "REQ-002" in req_ids
        assert "REQ-003" in req_ids

    def test_xlsx_meets_values(self):
        data = generate_xlsx(SAMPLE_DRAFT["responses"], SAMPLE_REQUIREMENTS)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        meets_values = [ws.cell(row=r, column=4).value for r in range(2, 5)]
        assert "Yes" in meets_values
        assert "No" in meets_values

    def test_xlsx_has_data_validation(self):
        data = generate_xlsx(SAMPLE_DRAFT["responses"], SAMPLE_REQUIREMENTS)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert len(ws.data_validations.dataValidation) >= 1

    def test_xlsx_handles_empty_responses(self):
        data = generate_xlsx([], [])
        assert isinstance(data, bytes)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.max_row == 1  # header only

    def test_xlsx_includes_source_urls(self):
        data = generate_xlsx(SAMPLE_DRAFT["responses"], SAMPLE_REQUIREMENTS)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        urls = [ws.cell(row=r, column=6).value for r in range(2, 5)]
        assert any(u and "azure.microsoft.com" in u for u in urls)
